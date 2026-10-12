"""The CEO assistant's own page: the briefing, decisions waiting, work handed out, alerts.

The same domain functions the assistant's tools call, so the page and the chat never
disagree. Using any of it needs "Dùng Trợ lý CEO"; changing the alert rules and the
digest schedule needs "Cấu hình cảnh báo điều hành". Every section still checks the
department's own permission for the person asking.
"""

from __future__ import annotations

import io
from datetime import date
from typing import Any, Literal, Optional

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.core.security import PermissionRequired
from app.domains.executive.alerts import evaluate_alerts, rule_settings, save_rule_settings, serialize_setting
from app.domains.executive.briefing import build_briefing
from app.domains.executive.digest import build_digest
from app.domains.executive.inbox import approvals_inbox
from app.domains.executive.scheduler import get_schedule, save_schedule, serialize_schedule
from app.domains.executive.tasks import (
    PRIORITY_LABELS,
    _describe,
    assign_task_to,
    find_people,
    list_tasks,
    person_in_reach,
)
from app.domains.hr.hr_employee_tools import directory_scope
from app.domains.platform.position_service import has_permission
from app.models.models import User

router = APIRouter(prefix="/executive", tags=["Executive assistant"])

USE = PermissionRequired("executive.briefing.view")
MANAGE = PermissionRequired("executive.alerts.manage")
MANAGE_PERMISSION = "executive.alerts.manage"
Period = Literal["TODAY", "WEEK", "MONTH", "QUARTER"]


@router.get("/briefing")
def get_briefing(
    period: Period = "TODAY",
    db: Session = Depends(get_db),
    user: User = Depends(USE),
) -> dict[str, Any]:
    briefing = build_briefing(db, user, period=period)
    db.commit()  # reading labour contracts leaves an access record
    return briefing


@router.get("/approvals")
def get_approvals(
    limit: int = Query(50, ge=1, le=200),
    db: Session = Depends(get_db),
    user: User = Depends(USE),
) -> dict[str, Any]:
    return approvals_inbox(db, user, limit=limit)


@router.get("/tasks")
def get_tasks(
    scope: Literal["ASSIGNED_BY_ME", "ASSIGNED_TO_ME"] = "ASSIGNED_BY_ME",
    status: Literal["OPEN", "OVERDUE", "DONE", "ALL"] = "OPEN",
    limit: int = Query(50, ge=1, le=200),
    db: Session = Depends(get_db),
    user: User = Depends(USE),
) -> dict[str, Any]:
    return list_tasks(db, user, scope=scope, status=status, limit=limit)


@router.get("/people")
def get_people(
    q: str = Query("", max_length=120),
    db: Session = Depends(get_db),
    user: User = Depends(USE),
) -> list[dict[str, Any]]:
    """Colleagues the user may hand work to: their directory, searched by name or email."""
    if q.strip():
        people = find_people(db, user, q)
    else:
        scope, _ = directory_scope(db, user)
        people = db.query(User).filter(
            User.tenant_id == user.tenant_id, User.is_active.is_(True), User.id.in_(scope),
        ).order_by(User.full_name).limit(200).all()
    return [
        {"id": str(person.id), "name": person.full_name, "label": _describe(person), "email": person.email}
        for person in people[:200]
    ]


class AssignRequest(BaseModel):
    assignee_id: str
    title: str = Field(min_length=1, max_length=255)
    description: Optional[str] = Field(None, max_length=5000)
    due_date: Optional[date] = None
    priority: Literal["LOW", "MEDIUM", "HIGH", "URGENT"] = "MEDIUM"


@router.post("/tasks", status_code=201)
def create_task(
    body: AssignRequest,
    db: Session = Depends(get_db),
    user: User = Depends(USE),
) -> dict[str, Any]:
    person = person_in_reach(db, user, body.assignee_id)
    if person is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy người nhận trong danh bạ bạn được xem")
    result = assign_task_to(
        db, user, person, title=body.title, description=body.description,
        due_date=body.due_date, priority=body.priority,
    )
    if not result["created"]:
        db.rollback()
        raise HTTPException(status_code=422, detail=result["reply"])
    db.commit()
    return result


def _settings_payload(db: Session, user: User) -> dict[str, Any]:
    schedule = get_schedule(db, user.tenant_id)
    return {
        "rules": [serialize_setting(setting) for setting in rule_settings(db, user.tenant_id)],
        "schedule": serialize_schedule(schedule) if schedule is not None else None,
        "can_manage": has_permission(db, user, MANAGE_PERMISSION),
    }


@router.get("/alerts")
def get_alerts(db: Session = Depends(get_db), user: User = Depends(USE)) -> dict[str, Any]:
    """What the rules find right now for this person, and how they are set."""
    alerts = evaluate_alerts(db, user)
    payload = _settings_payload(db, user)
    db.commit()  # the schedule row is created on first read; contract reads are recorded
    return {"alerts": alerts, **payload}


class RuleChange(BaseModel):
    key: str
    enabled: bool = True
    threshold: float


class ScheduleChange(BaseModel):
    digest_enabled: Optional[bool] = None
    digest_hour: Optional[int] = None
    digest_weekdays: Optional[list[int]] = None
    alerts_enabled: Optional[bool] = None


class AlertSettingsRequest(BaseModel):
    rules: list[RuleChange] = Field(default_factory=list, max_length=20)
    schedule: Optional[ScheduleChange] = None


@router.put("/alert-settings")
def put_alert_settings(
    body: AlertSettingsRequest,
    db: Session = Depends(get_db),
    user: User = Depends(MANAGE),
) -> dict[str, Any]:
    if body.rules:
        save_rule_settings(db, user, [change.model_dump() for change in body.rules])
    if body.schedule is not None:
        save_schedule(db, user, body.schedule.model_dump(exclude_none=True))
    db.commit()
    return _settings_payload(db, user)


@router.get("/digest")
def get_digest(db: Session = Depends(get_db), user: User = Depends(USE)) -> dict[str, Any]:
    """The morning digest as this person would receive it now."""
    digest = build_digest(db, user, len(evaluate_alerts(db, user)))
    db.commit()
    return digest


_SECTION_TITLES = {
    "finance": "Tài chính",
    "people": "Nhân sự",
    "legal": "Pháp chế",
    "marketing": "Marketing",
    "approvals": "Chờ duyệt",
    "tasks": "Công việc",
    "ai_cost": "Chi phí AI",
}


def _flatten(value: Any, prefix: str = "") -> list[tuple[str, str]]:
    """Leaf values of a section as (path, value) rows, lists numbered."""
    rows: list[tuple[str, str]] = []
    if isinstance(value, dict):
        for key, item in value.items():
            if key in {"source", "page", "note", "permitted", "link"}:
                continue
            rows += _flatten(item, f"{prefix}.{key}" if prefix else key)
    elif isinstance(value, list):
        for index, item in enumerate(value, start=1):
            rows += _flatten(item, f"{prefix}[{index}]")
    elif value is not None:
        rows.append((prefix, str(value)))
    return rows


@router.get("/briefing/export")
def export_briefing(
    period: Period = "MONTH",
    db: Session = Depends(get_db),
    user: User = Depends(USE),
):
    """The briefing as an Excel workbook: the digest, then one sheet per section."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    briefing = build_briefing(db, user, period=period)
    alerts = evaluate_alerts(db, user)
    digest = build_digest(db, user, len(alerts))
    db.commit()
    book = Workbook()
    bold, head = Font(bold=True), PatternFill("solid", fgColor="DCE6F2")

    summary = book.active
    summary.title = "Tóm tắt"
    summary.append([f"Bản tin điều hành ({briefing['period']})", f"{briefing['from']} → {briefing['to']}"])
    summary["A1"].font = Font(bold=True, size=13)
    summary.append([])
    for line in digest["lines"]:
        summary.append([line])
    summary.append([])
    summary.append(["Cảnh báo", "Mức", "Chi tiết"])
    for cell in summary[summary.max_row]:
        cell.font, cell.fill = bold, head
    for alert in alerts:
        summary.append([alert["title"], "Nghiêm trọng" if alert["severity"] == "CRITICAL" else "Cần chú ý", alert["message"]])
    summary.column_dimensions["A"].width = 60
    summary.column_dimensions["C"].width = 90

    for name, section in briefing["sections"].items():
        sheet = book.create_sheet(_SECTION_TITLES.get(name, name)[:31])
        sheet.append(["Chỉ tiêu", "Giá trị"])
        for cell in sheet[1]:
            cell.font, cell.fill = bold, head
        if isinstance(section, dict) and section.get("permitted") is False:
            sheet.append(["Không có quyền xem", ", ".join(section.get("needs") or [])])
        else:
            for path, value in _flatten(section):
                sheet.append([path, value])
        sheet.column_dimensions["A"].width = 55
        sheet.column_dimensions["B"].width = 60
        for row in sheet.iter_rows(min_row=2):
            row[1].alignment = Alignment(wrap_text=True, vertical="top")

    buffer = io.BytesIO()
    book.save(buffer)
    buffer.seek(0)
    filename = f"ban-tin-dieu-hanh-{briefing['from']}-{briefing['to']}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/meta")
def get_meta(db: Session = Depends(get_db), user: User = Depends(USE)) -> dict[str, Any]:
    return {"priorities": PRIORITY_LABELS, "can_manage_alerts": has_permission(db, user, MANAGE_PERMISSION)}
