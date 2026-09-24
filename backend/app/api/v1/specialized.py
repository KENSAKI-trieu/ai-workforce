"""
API Endpoints for Legal, IT, Finance, and Sales domain operations.
"""

import io
import json
import mimetypes
import uuid
from pathlib import Path
from typing import Dict, Any, Optional, List
from urllib.parse import quote

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import PlainTextResponse, Response
from pydantic import BaseModel
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.core.security import get_current_active_user
from app.models.models import AIAgent, AgentWorkflow, ContractReview, User, WorkflowApproval
from app.core.tool_permissions import grant_decision
from app.plugins.resolver import resolve_skill_restriction
from app.services.audit_service import log_audit_action
from app.services import contract_review_store
from app.services.document_parser import DocumentParseError, extract_file_text
from app.services.legal_service import (
    audit_contract_text,
    check_software_licenses,
    compare_contract_texts,
    detect_sensitive_data,
)
from app.services.contract_review import detect_contract_type, review_contract
from app.services.contract_redline import build_redline_docx
# Shared with the chat path so both entry points escalate identically.
from app.services.legal_approval_service import create_legal_approval as _create_legal_approval
from app.services.legal_document_generator import generate_legal_document
from app.services.legal_draft_storage import read_legal_artifact, save_legal_artifact
from app.services.legal_documents import list_document_schemas, validate_document_fields
from app.services.notification_service import create_notification
from app.services.rag_service import hybrid_search_documents
from app.core.agent_status import refuse_under_development
from app.domains.incubating.finance_service import audit_invoice_and_reconcile
from app.domains.incubating.it_service import handle_it_request
from app.domains.incubating.sales_service import handle_sales_request

router = APIRouter(tags=["Specialized Domain APIs"])


def _legal_tool_required(tool_name: str):
    """Refuse a Legal action whose tool is switched off for the tenant's Legal agent.

    These endpoints used to run whatever the configuration page said: switching a Legal
    tool off there changed nothing, so an operator who revoked it had no way to know it
    was still in use.
    """

    def dependency(
        db: Session = Depends(get_db),
        current_user: User = Depends(get_current_active_user),
    ) -> None:
        agent = db.query(AIAgent).filter(
            AIAgent.tenant_id == current_user.tenant_id,
            AIAgent.role_code == "LEGAL",
        ).first()
        # Plugin narrowing counts here too: these endpoints used to check the agent's
        # own grants only, so a tenant package withdrawing a Legal tool left its page
        # working.
        if agent is None or grant_decision(
            tool_name,
            tools_access=agent.tools_access,
            allowed_actions=agent.allowed_actions,
            disallowed_actions=agent.disallowed_actions,
            restriction=resolve_skill_restriction(db, current_user.tenant_id, "LEGAL"),
        ) is not None:
            raise HTTPException(
                status_code=403,
                detail=(
                    f"Công cụ `{tool_name}` chưa được bật cho Legal Counsel AI. "
                    "Admin hoặc Owner có thể bật trong phần Cấu hình AI Employees."
                ),
            )

    return Depends(dependency)
MAX_LEGAL_FILE_BYTES = 10 * 1024 * 1024
LEGAL_DOCUMENT_APPROVERS = {"Owner", "Admin", "CEO"}


class ContractAuditRequest(BaseModel):
    document_name: str = "Contract.pdf"
    contract_text: str
    represented_party: str = "NEUTRAL"


class CreateJiraTicketRequest(BaseModel):
    summary: str
    description: Optional[str] = None
    priority: str = "MEDIUM"


class AuditInvoiceRequest(BaseModel):
    invoice_text: str


class SalesQuotationRequest(BaseModel):
    customer_name: str = "Khách Hàng Doanh Nghiệp"
    item_query: str


class LegalDocumentGenerateRequest(BaseModel):
    document_type: str
    output_format: str = "docx"
    fields: dict[str, Any]


class LegalDocumentValidationRequest(BaseModel):
    document_type: str
    fields: dict[str, Any]


class ContractReviewDecisionRequest(BaseModel):
    decision: str
    revised_text: Optional[str] = None
    comment: Optional[str] = None


def _legal_draft_approval(
    db: Session,
    current_user: User,
    artifact_id: str,
) -> WorkflowApproval:
    approvals = db.query(WorkflowApproval).join(AgentWorkflow).filter(
        AgentWorkflow.tenant_id == current_user.tenant_id,
        WorkflowApproval.action_type == "LEGAL_DOCUMENT_APPROVAL",
    ).all()
    approval = next(
        (
            item
            for item in approvals
            if str((item.payload or {}).get("artifact_id")) == artifact_id
        ),
        None,
    )
    if not approval:
        raise HTTPException(status_code=404, detail="Legal document draft not found")
    return approval


def _legal_draft_item(approval: WorkflowApproval, current_user: User) -> dict[str, Any]:
    payload = approval.payload or {}
    is_reviewer = current_user.role in LEGAL_DOCUMENT_APPROVERS
    is_creator = approval.workflow.initiator_id == current_user.id
    approved = approval.status == "APPROVED"
    download_variant = "draft" if is_reviewer and not approved else "approved"
    artifact_id = str(payload.get("artifact_id"))
    return {
        "artifact_id": artifact_id,
        "approval_id": str(approval.id),
        "workflow_id": str(approval.workflow_id),
        "document_type": payload.get("document_type"),
        "document_type_label": payload.get("document_type_label"),
        "filename": payload.get("filename"),
        "output_format": payload.get("output_format"),
        "status": approval.status,
        "comments": approval.comments,
        "requester_name": payload.get("requester_name"),
        "requester_id": payload.get("requester_id"),
        "submitted_at": (
            approval.workflow.created_at.isoformat()
            if approval.workflow.created_at
            else None
        ),
        "updated_at": approval.updated_at.isoformat() if approval.updated_at else None,
        "can_preview": is_reviewer or is_creator,
        "can_download": is_reviewer or (is_creator and approved),
        "download_variant": download_variant,
        "preview_url": f"/api/v1/legal/document-drafts/{artifact_id}/preview",
        "download_url": f"/api/v1/legal/document-drafts/{artifact_id}/download?variant={download_variant}",
    }


def _contract_review_for_user(
    db: Session, current_user: User, review_id: str
) -> ContractReview:
    review = contract_review_store.get_contract_review(
        db, user=current_user, review_id=review_id
    )
    if not review:
        # A review in another tenant, and one this user may not read, are the same
        # 404: the existence of a contract is itself information.
        raise HTTPException(status_code=404, detail="Contract review not found")
    return review


def _contract_review_item(
    review: ContractReview, decisions: list[dict[str, Any]], current_user: User
) -> dict[str, Any]:
    accepted = sum(item["decision"] in {"ACCEPTED", "EDITED"} for item in decisions)
    return {
        "review_id": str(review.id),
        "workflow_id": str(review.workflow_id) if review.workflow_id else None,
        "document_name": review.document_name,
        "source": review.source,
        "contract_type": review.contract_type,
        "contract_type_label": (review.result or {}).get("contract_type_label"),
        "represented_party": review.represented_party,
        "represented_party_label": (review.result or {}).get("represented_party_label"),
        "risk_score": review.risk_score,
        "risk_level": review.risk_level,
        "total_findings": review.total_findings,
        "decided_count": len(decisions),
        "accepted_count": accepted,
        "status": review.status,
        "created_by_name": review.created_by.full_name if review.created_by else None,
        "created_at": review.created_at.isoformat() if review.created_at else None,
        "updated_at": review.updated_at.isoformat() if review.updated_at else None,
        "redline_ready": accepted > 0,
        "redline_url": f"/api/v1/legal/contract-reviews/{review.id}/redline",
        "can_decide": contract_review_store.can_access_contract_review(current_user, review),
    }


async def _read_legal_file(file: UploadFile) -> tuple[str, str, list[str]]:
    data = await file.read()
    if not data:
        raise HTTPException(status_code=422, detail="The uploaded file is empty")
    if len(data) > MAX_LEGAL_FILE_BYTES:
        raise HTTPException(status_code=413, detail="Legal files are limited to 10 MB")
    filename = Path(file.filename or "document.txt").name
    extension = Path(filename).suffix.lower()
    headers: list[str] = []
    try:
        if extension == ".json":
            payload = json.loads(data.decode("utf-8-sig"))
            text = json.dumps(payload, ensure_ascii=False, indent=2)
            headers = list(payload) if isinstance(payload, dict) else []
        elif extension == ".xlsx":
            from openpyxl import load_workbook

            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            rows: list[str] = []
            for sheet in workbook.worksheets:
                rows.append(f"# {sheet.title}")
                for row_index, row in enumerate(sheet.iter_rows(values_only=True)):
                    values = ["" if value is None else str(value) for value in row]
                    if row_index == 0:
                        headers.extend(value for value in values if value)
                    rows.append(" | ".join(values))
            text = "\n".join(rows)
        else:
            text = extract_file_text(filename, data)
            if extension == ".csv" and text:
                headers = [part.strip() for part in text.splitlines()[0].split("|")]
    except (DocumentParseError, UnicodeDecodeError, ValueError, json.JSONDecodeError) as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    if not text.strip():
        raise HTTPException(status_code=422, detail="No readable text was found in the file")
    return filename, text, headers


# --- LEGAL ---
def _retrieve_contract_review_references(
    db: Session,
    current_user: User,
    contract_type_label: str,
) -> list[dict[str, Any]]:
    """Retrieve ACL-filtered internal policy/template context for Legal review."""
    try:
        chunks = hybrid_search_documents(
            db=db,
            tenant_id=current_user.tenant_id,
            query_text=(
                f"{contract_type_label} mẫu hợp đồng chuẩn policy pháp lý "
                "thanh toán trách nhiệm sở hữu trí tuệ chấm dứt bảo mật"
            ),
            department=(
                "*"
                if current_user.role in {"Owner", "Admin", "CEO"}
                else current_user.department
            ),
            top_k=6,
            user_role=current_user.role,
            user_department=current_user.department,
        )
    except Exception:
        # Contract analysis still works deterministically if the tenant has no
        # indexed legal knowledge or its retrieval service is temporarily down.
        db.rollback()
        return []

    references: list[dict[str, Any]] = []
    seen_documents: set[tuple[str, str]] = set()
    for chunk in chunks:
        name = str(chunk.get("document_title") or chunk.get("document_name") or "Tài liệu nội bộ")
        document_id = str(chunk.get("document_id") or chunk.get("document_name") or "")
        version = str(chunk.get("version") or "1.0")
        document_key = (document_id, version)
        if not document_id or document_key in seen_documents:
            continue
        seen_documents.add(document_key)
        normalized_name = name.lower()
        source_type = (
            "APPROVED_TEMPLATE"
            if any(token in normalized_name for token in ("template", "mẫu", "mau", "approved", "chuẩn"))
            else "COMPANY_POLICY"
        )
        references.append({
            "id": document_id,
            "document_id": document_id,
            "version": version,
            "type": source_type,
            "title": name,
            "section_title": chunk.get("section_title"),
            "citation_tag": chunk.get("citation_tag"),
            "score": chunk.get("score"),
            "url": "",
            "reader_url": (
                f"/api/v1/documents/{quote(document_id, safe='')}/reader"
                f"?version={quote(version, safe='')}"
            ),
            "note": "Ngữ cảnh nội bộ đã truy xuất theo ACL; Legal cần xác nhận tính áp dụng.",
        })
    return references


@router.get("/legal/document-templates", summary="List schema-driven legal document templates")
def list_legal_document_templates(
    current_user: User = Depends(get_current_active_user),
) -> List[Dict[str, Any]]:
    return list_document_schemas()


@router.post("/legal/validate-document", summary="Validate a legal document draft before generation")
def validate_legal_document_endpoint(
    req: LegalDocumentValidationRequest,
    current_user: User = Depends(get_current_active_user),
) -> Dict[str, Any]:
    try:
        return validate_document_fields(req.document_type, req.fields)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


@router.post(
    "/legal/audit-contract",
    summary="Audit contract text for high-risk clauses",
    dependencies=[_legal_tool_required("audit_contract_risk")],
)
def audit_contract_endpoint(
    req: ContractAuditRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user),
) -> Dict[str, Any]:
    try:
        result = review_contract(
            req.contract_text,
            req.document_name,
            req.represented_party,
        )
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    review = contract_review_store.save_contract_review(
        db,
        user=current_user,
        result=result,
        contract_text=req.contract_text,
        source="API",
    )
    db.commit()
    result["review_id"] = str(review.id)
    result["redline_url"] = f"/api/v1/legal/contract-reviews/{review.id}/redline"
    return result


@router.post(
    "/legal/review-document",
    summary="Extract and review a legal document",
    dependencies=[_legal_tool_required("audit_contract_risk")],
)
async def review_legal_document(
    file: UploadFile = File(...),
    represented_party: str = Form(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user),
) -> Dict[str, Any]:
    filename, text, _ = await _read_legal_file(file)
    detection = detect_contract_type(text)
    references = _retrieve_contract_review_references(
        db, current_user, detection["contract_type_label"]
    )
    try:
        result = review_contract(text, filename, represented_party, references)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    # Saved here rather than by a follow-up call from the browser: a second call
    # would have to accept the review body back from the client, which would let
    # anyone post a fabricated result and have it become the audit record.
    #
    # The review is stored before the escalation so its id can be handed to
    # create_legal_approval, which uses it to reuse the approval already waiting on
    # this same review instead of opening a second card for one contract. The
    # approval then carries the id too, so the approvals screen can open the review.
    review = contract_review_store.save_contract_review(
        db,
        user=current_user,
        result=result,
        contract_text=text,
        source="UPLOAD",
    )
    result["workflow_id"] = _create_legal_approval(
        db, current_user, result, contract_review_id=str(review.id)
    )
    result["approval_created"] = result["workflow_id"] is not None
    if result["workflow_id"]:
        review.workflow_id = uuid.UUID(result["workflow_id"])
        # The stored blob is what a reopened review renders from, so it has to carry
        # the escalation as well as the analyzer output.
        review.result = {
            **(review.result or {}),
            "workflow_id": result["workflow_id"],
            "approval_created": True,
        }
    db.commit()
    result["review_id"] = str(review.id)
    result["decisions"] = contract_review_store.serialize_decisions(db, review)
    result["redline_url"] = f"/api/v1/legal/contract-reviews/{review.id}/redline"
    return result


@router.get("/legal/contract-reviews", summary="List saved contract reviews")
def list_contract_reviews_endpoint(
    limit: int = 50,
    offset: int = 0,
    risk_level: Optional[str] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user),
) -> list[dict[str, Any]]:
    reviews = contract_review_store.list_contract_reviews(
        db, user=current_user, limit=limit, offset=offset, risk_level=risk_level
    )
    return [
        _contract_review_item(
            review, contract_review_store.serialize_decisions(db, review), current_user
        )
        for review in reviews
    ]


@router.get(
    "/legal/contract-reviews/{review_id}",
    summary="Reopen a saved contract review with its decisions",
)
def get_contract_review_endpoint(
    review_id: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user),
) -> Dict[str, Any]:
    review = _contract_review_for_user(db, current_user, review_id)
    decisions = contract_review_store.serialize_decisions(db, review)
    # The stored analyzer output is spread at the top level so the client renders a
    # reopened review through exactly the same shape as a fresh one.
    return {
        **(review.result or {}),
        **_contract_review_item(review, decisions, current_user),
        "decisions": decisions,
    }


@router.put(
    "/legal/contract-reviews/{review_id}/decisions/{finding_key}",
    summary="Record the reviewer decision for one finding",
)
def put_contract_review_decision(
    review_id: str,
    finding_key: str,
    req: ContractReviewDecisionRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user),
) -> Dict[str, Any]:
    review = _contract_review_for_user(db, current_user, review_id)
    decision = (req.decision or "").upper()
    if decision not in contract_review_store.VALID_DECISIONS:
        raise HTTPException(
            status_code=422, detail="decision must be ACCEPTED, REJECTED or EDITED"
        )
    if finding_key not in contract_review_store.finding_keys(review):
        raise HTTPException(
            status_code=422, detail="This finding does not belong to the review"
        )
    if decision == "EDITED" and not (req.revised_text or "").strip():
        raise HTTPException(
            status_code=422, detail="revised_text is required when the decision is EDITED"
        )
    contract_review_store.upsert_decision(
        db,
        review=review,
        user=current_user,
        finding_key=finding_key,
        decision=decision,
        revised_text=req.revised_text,
        comment=req.comment,
    )
    decisions = contract_review_store.serialize_decisions(db, review)
    payload = {
        **_contract_review_item(review, decisions, current_user),
        "decisions": decisions,
    }
    # log_audit_action commits internally, so it has to be the last write: anything
    # after it would land in a separate transaction.
    log_audit_action(
        db,
        current_user.tenant_id,
        "LEGAL",
        "record_contract_review_decision",
        {"review_id": str(review.id), "finding_key": finding_key, "decision": decision},
        {"decided_count": len(decisions), "status": review.status},
    )
    return payload


@router.delete(
    "/legal/contract-reviews/{review_id}/decisions/{finding_key}",
    summary="Clear the reviewer decision for one finding",
)
def delete_contract_review_decision(
    review_id: str,
    finding_key: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user),
) -> Dict[str, Any]:
    review = _contract_review_for_user(db, current_user, review_id)
    if not contract_review_store.clear_decision(db, review=review, finding_key=finding_key):
        raise HTTPException(status_code=404, detail="No decision recorded for this finding")
    decisions = contract_review_store.serialize_decisions(db, review)
    payload = {
        **_contract_review_item(review, decisions, current_user),
        "decisions": decisions,
    }
    log_audit_action(
        db,
        current_user.tenant_id,
        "LEGAL",
        "clear_contract_review_decision",
        {"review_id": str(review.id), "finding_key": finding_key},
        {"decided_count": len(decisions), "status": review.status},
    )
    return payload


@router.post(
    "/legal/compare-documents",
    summary="Compare two contract versions",
    dependencies=[_legal_tool_required("compare_contract_versions")],
)
async def compare_legal_documents(
    old_file: UploadFile = File(...),
    new_file: UploadFile = File(...),
    current_user: User = Depends(get_current_active_user),
) -> Dict[str, Any]:
    old_name, old_text, _ = await _read_legal_file(old_file)
    new_name, new_text, _ = await _read_legal_file(new_file)
    result = compare_contract_texts(old_text, new_text)
    return {"old_document": old_name, "new_document": new_name, **result}


@router.post(
    "/legal/privacy-check",
    summary="Detect personal and restricted data",
    dependencies=[_legal_tool_required("check_sensitive_data")],
)
async def privacy_check_document(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user),
) -> Dict[str, Any]:
    filename, text, headers = await _read_legal_file(file)
    result = {"document_name": filename, **detect_sensitive_data(text, headers)}
    result["workflow_id"] = _create_legal_approval(
        db, current_user, result, "LEGAL_PRIVACY_APPROVAL"
    )
    result["approval_created"] = result["workflow_id"] is not None
    return result


@router.post(
    "/legal/license-check",
    summary="Inspect a software dependency manifest",
    dependencies=[_legal_tool_required("check_software_licenses")],
)
async def license_check_manifest(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user),
) -> Dict[str, Any]:
    filename, text, _ = await _read_legal_file(file)
    result = check_software_licenses(filename, text)
    result["workflow_id"] = _create_legal_approval(
        db, current_user, result, "LEGAL_LICENSE_APPROVAL"
    )
    result["approval_created"] = result["workflow_id"] is not None
    return result


@router.post("/legal/generate-document", summary="Generate an editable legal document draft")
def generate_legal_document_endpoint(
    req: LegalDocumentGenerateRequest,
    current_user: User = Depends(get_current_active_user),
):
    raise HTTPException(
        status_code=410,
        detail="Direct generation is disabled; submit through /legal/document-drafts",
    )


@router.post(
    "/legal/document-drafts",
    status_code=201,
    summary="Generate and submit a legal document for approval",
    dependencies=[_legal_tool_required("generate_legal_document")],
)
def submit_legal_document_draft(
    req: LegalDocumentGenerateRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user),
) -> dict[str, Any]:
    try:
        draft_content, filename, media_type = generate_legal_document(
            req.document_type, req.output_format, req.fields, approved=False
        )
        approved_content, _, _ = generate_legal_document(
            req.document_type, req.output_format, req.fields, approved=True
        )
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc

    artifact_id = uuid.uuid4().hex
    draft_key = save_legal_artifact(
        tenant_id=current_user.tenant_id,
        artifact_id=artifact_id,
        variant="draft",
        filename=filename,
        content=draft_content,
    )
    approved_key = save_legal_artifact(
        tenant_id=current_user.tenant_id,
        artifact_id=artifact_id,
        variant="approved",
        filename=filename,
        content=approved_content,
    )
    template = next(
        (
            item
            for item in list_document_schemas()
            if item["id"] == req.document_type.upper()
        ),
        None,
    )
    document_type_label = template["label"] if template else req.document_type
    workflow = AgentWorkflow(
        tenant_id=current_user.tenant_id,
        initiator_id=current_user.id,
        title=f"Phê duyệt văn bản: {document_type_label}",
        status="AWAITING_APPROVAL",
        current_step=1,
        dag_plan={
            "agent_role": "LEGAL",
            "steps": ["DRAFT_CREATED", "EXECUTIVE_APPROVAL", "CREATOR_DOWNLOAD"],
        },
    )
    db.add(workflow)
    db.flush()
    approval = WorkflowApproval(
        workflow_id=workflow.id,
        action_type="LEGAL_DOCUMENT_APPROVAL",
        risk_level="MEDIUM",
        payload={
            "artifact_id": artifact_id,
            "document_type": req.document_type.upper(),
            "document_type_label": document_type_label,
            "filename": filename,
            "media_type": media_type,
            "output_format": req.output_format.lower(),
            "draft_storage_key": draft_key,
            "approved_storage_key": approved_key,
            "requester_id": str(current_user.id),
            "requester_name": current_user.full_name,
            "reason": "Văn bản do Legal Agent tạo cần CEO, Admin hoặc Owner phê duyệt trước khi người tạo tải xuống.",
            "data_sources": [document_type_label],
            "preview_url": f"/api/v1/legal/document-drafts/{artifact_id}/preview",
            "review_download_url": f"/api/v1/legal/document-drafts/{artifact_id}/download?variant=draft",
        },
        status="WAITING",
    )
    db.add(approval)
    for approver in db.query(User).filter(
        User.tenant_id == current_user.tenant_id,
        User.role.in_(LEGAL_DOCUMENT_APPROVERS),
        User.is_active.is_(True),
    ).all():
        create_notification(
            db,
            user=approver,
            event_type="APPROVAL_REQUIRED",
            title="Văn bản pháp lý chờ phê duyệt",
            message=f"{current_user.full_name} đã gửi {document_type_label}.",
            severity="WARNING",
            entity_type="APPROVAL",
            entity_id=str(approval.id),
            dedup_key=f"legal-document-approval:{approval.id}:{approver.id}",
        )
    db.commit()
    db.refresh(approval)
    return _legal_draft_item(approval, current_user)


@router.get("/legal/document-drafts", summary="List generated legal documents")
def list_legal_document_drafts(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user),
) -> list[dict[str, Any]]:
    query = db.query(WorkflowApproval).join(AgentWorkflow).filter(
        AgentWorkflow.tenant_id == current_user.tenant_id,
        WorkflowApproval.action_type == "LEGAL_DOCUMENT_APPROVAL",
    )
    if current_user.role not in LEGAL_DOCUMENT_APPROVERS:
        query = query.filter(AgentWorkflow.initiator_id == current_user.id)
    approvals = query.order_by(AgentWorkflow.created_at.desc()).all()
    return [_legal_draft_item(approval, current_user) for approval in approvals]


@router.get(
    "/legal/document-drafts/{artifact_id}/preview",
    summary="Preview a generated legal document",
)
def preview_legal_document_draft(
    artifact_id: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user),
) -> dict[str, Any]:
    approval = _legal_draft_approval(db, current_user, artifact_id)
    if (
        current_user.role not in LEGAL_DOCUMENT_APPROVERS
        and approval.workflow.initiator_id != current_user.id
    ):
        raise HTTPException(status_code=403, detail="You cannot preview this legal document")
    payload = approval.payload or {}
    try:
        content = read_legal_artifact(str(payload["draft_storage_key"]))
        text = extract_file_text(str(payload["filename"]), content)
    except (KeyError, OSError, ValueError, DocumentParseError) as exc:
        raise HTTPException(status_code=404, detail="Legal document artifact is unavailable") from exc
    return {
        "artifact_id": artifact_id,
        "filename": payload.get("filename"),
        "document_type_label": payload.get("document_type_label"),
        "status": approval.status,
        "requester_name": payload.get("requester_name"),
        "content": text,
    }


@router.get(
    "/legal/document-drafts/{artifact_id}/download",
    summary="Download a generated legal document with approval enforcement",
)
def download_legal_document_draft(
    artifact_id: str,
    variant: str = "approved",
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user),
) -> Response:
    approval = _legal_draft_approval(db, current_user, artifact_id)
    is_reviewer = current_user.role in LEGAL_DOCUMENT_APPROVERS
    is_creator = approval.workflow.initiator_id == current_user.id
    if variant not in {"draft", "approved"}:
        raise HTTPException(status_code=422, detail="Unsupported document variant")
    if variant == "draft" and not is_reviewer:
        raise HTTPException(status_code=403, detail="Only executive approvers can download the draft")
    if variant == "approved" and approval.status != "APPROVED":
        raise HTTPException(
            status_code=403,
            detail="The approved artifact is unavailable until approval is completed",
        )
    if variant == "approved" and not is_reviewer and not is_creator:
        raise HTTPException(
            status_code=403,
            detail="Only the creator or an executive approver can download this document",
        )
    payload = approval.payload or {}
    storage_field = "draft_storage_key" if variant == "draft" else "approved_storage_key"
    try:
        content = read_legal_artifact(str(payload[storage_field]))
    except (KeyError, OSError, ValueError) as exc:
        raise HTTPException(status_code=404, detail="Legal document artifact is unavailable") from exc
    filename = Path(str(payload.get("filename") or "legal-document")).name
    media_type = str(payload.get("media_type") or mimetypes.guess_type(filename)[0] or "application/octet-stream")
    disposition_name = filename if variant == "draft" else f"approved-{filename}"
    return Response(
        content=content,
        media_type=media_type,
        headers={
            "Content-Disposition": (
                f"attachment; filename=\"{disposition_name}\"; "
                f"filename*=UTF-8''{quote(disposition_name)}"
            )
        },
    )


@router.get(
    "/legal/contract-reviews/{review_id}/redline",
    summary="Download the redline report built from the accepted revisions",
)
def download_contract_redline(
    review_id: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user),
) -> Response:
    review = _contract_review_for_user(db, current_user, review_id)
    decisions = contract_review_store.serialize_decisions(db, review)
    if not any(item["decision"] in {"ACCEPTED", "EDITED"} for item in decisions):
        raise HTTPException(
            status_code=409,
            detail="Chưa có đề xuất nào được chấp nhận để tạo redline.",
        )

    content: bytes | None = None
    if review.redline_storage_key:
        try:
            content = read_legal_artifact(str(review.redline_storage_key))
        except (OSError, ValueError):
            # The cached file is gone; fall through and rebuild it.
            content = None
    if content is None:
        content, filename = build_redline_docx(
            review.result or {},
            decisions,
            document_name=review.document_name,
            generated_by=current_user.full_name,
            review_id=str(review.id),
        )
        artifact_id = review.redline_artifact_id or uuid.uuid4().hex
        review.redline_artifact_id = artifact_id
        review.redline_filename = filename
        review.redline_storage_key = save_legal_artifact(
            tenant_id=review.tenant_id,
            artifact_id=artifact_id,
            variant="redline",
            filename=filename,
            content=content,
        )
        db.commit()

    filename = Path(str(review.redline_filename or "redline.docx")).name
    return Response(
        content=content,
        media_type=(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ),
        headers={
            "Content-Disposition": (
                f"attachment; filename=\"{filename}\"; "
                f"filename*=UTF-8''{quote(filename)}"
            )
        },
    )


# --- IT ---
@router.post("/it/tickets", summary="Create Jira ticket for technical issue")
def create_jira_ticket_endpoint(
    req: CreateJiraTicketRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user),
) -> Dict[str, Any]:
    # The IT agent is under development: this issued a Jira key no Jira ever saw.
    refuse_under_development("IT")
    return handle_it_request(db, current_user, f"{req.summary} - {req.description or ''}")


# --- FINANCE ---
@router.post("/finance/audit-invoice", summary="Audit invoice text and reconcile PO")
def audit_invoice_endpoint(
    req: AuditInvoiceRequest,
    current_user: User = Depends(get_current_active_user),
) -> Dict[str, Any]:
    # The Finance agent is under development: every invoice was checked against one fixed PO.
    refuse_under_development("FINANCE")
    return audit_invoice_and_reconcile(req.invoice_text)


# --- SALES ---
@router.post("/sales/quotation", summary="Generate sales quotation PDF payload")
def generate_quotation_endpoint(
    req: SalesQuotationRequest,
    current_user: User = Depends(get_current_active_user),
) -> Dict[str, Any]:
    # The Sales agent is under development: every request was quoted the same camera.
    refuse_under_development("SALES")
    return handle_sales_request(req.item_query, customer_name=req.customer_name)


@router.get("/sales/download-quote/{file_id}", response_class=PlainTextResponse, summary="Download sales PDF quotation file")
def download_quote(file_id: str):
    refuse_under_development("SALES")
    return f"SIMULATED PDF QUOTATION FILE FOR {file_id}\nOfficial AI Workforce Quotation Document."
