"""Bringing a company's books in from the Excel its accounting software exports.

Each kind of import has a fixed set of columns, matched by header text (Vietnamese label
or English key, accents and case ignored) so a MISA or Fast export only needs its header
row renamed. The whole file is checked before anything is written: one bad row rejects
the import with every problem listed by row and column, because half an import leaves
books that no longer balance and nobody can tell which half arrived.
"""

from __future__ import annotations

import io
import re
import unicodedata
import uuid
from collections import defaultdict
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from typing import Any, Callable

from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session

from app.domains.finance.charts import normal_balance, parent_code
from app.domains.finance.money import (
    ZERO,
    normalize_period,
    normalize_tax_code,
    period_of,
    to_date,
    to_decimal,
)
from app.models.models import (
    FinAccount,
    FinBudget,
    FinImportBatch,
    FinInvoice,
    FinJournalEntry,
    FinLedgerLine,
    FinParty,
    FinPayment,
    FinPurchaseOrder,
    User,
)

MAX_ROWS = 50_000


@dataclass(frozen=True)
class Column:
    key: str
    label: str
    required: bool = False


@dataclass(frozen=True)
class ImportKind:
    title: str
    columns: tuple[Column, ...]
    guide: str
    # Whether a batch can be undone: rows that only this import created can be removed;
    # upserts into shared master data (accounts, parties) cannot be told apart afterwards.
    undoable: bool = True


KINDS: dict[str, ImportKind] = {
    "accounts": ImportKind(
        "Hệ thống tài khoản",
        (Column("code", "Số TK", True), Column("name", "Tên TK", True)),
        "Thêm hoặc đổi tên tài khoản. TK cha tự suy từ số TK (1331 thuộc 133).",
        undoable=False,
    ),
    "parties": ImportKind(
        "Đối tượng (nhà cung cấp, khách hàng)",
        (
            Column("kind", "Loại (NCC/KH/CẢ HAI)", True),
            Column("tax_code", "MST"),
            Column("name", "Tên", True),
            Column("address", "Địa chỉ"),
            Column("email", "Email"),
            Column("bank_account", "Số TK ngân hàng"),
            Column("bank_name", "Ngân hàng"),
            Column("payment_terms_days", "Số ngày được nợ"),
        ),
        "Đối tượng trùng MST được cập nhật. Đối tượng không có MST luôn được thêm mới.",
        undoable=False,
    ),
    "opening_balances": ImportKind(
        "Số dư đầu kỳ",
        (
            Column("period", "Kỳ (YYYY-MM)", True),
            Column("account_code", "Số TK", True),
            Column("debit", "Dư Nợ"),
            Column("credit", "Dư Có"),
            Column("party_tax_code", "MST đối tượng"),
            Column("department", "Phòng ban"),
        ),
        "Tổng Dư Nợ phải bằng tổng Dư Có. Với TK 131/331, ghi chi tiết từng đối tượng "
        "thay cho một dòng tổng, để không bị cộng hai lần.",
    ),
    "ledger": ImportKind(
        "Sổ nhật ký chung / sổ cái",
        (
            Column("entry_date", "Ngày hạch toán", True),
            Column("voucher_no", "Số chứng từ"),
            Column("description", "Diễn giải"),
            Column("account_code", "Số TK", True),
            Column("debit", "Phát sinh Nợ"),
            Column("credit", "Phát sinh Có"),
            Column("party_tax_code", "MST đối tượng"),
            Column("department", "Phòng ban"),
        ),
        "Mỗi dòng một vế. Mỗi số chứng từ phải cân Nợ = Có, và cả file cũng vậy.",
    ),
    "budgets": ImportKind(
        "Ngân sách",
        (
            Column("department", "Phòng ban", True),
            Column("account_code", "Số TK", True),
            Column("period", "Kỳ (YYYY-MM)", True),
            Column("amount", "Số tiền", True),
        ),
        "Dòng trùng phòng ban + TK + kỳ với ngân sách đã có sẽ thay số cũ.",
    ),
    "purchase_orders": ImportKind(
        "Đơn mua hàng (PO)",
        (
            Column("po_number", "Số PO", True),
            Column("party_tax_code", "MST nhà cung cấp", True),
            Column("order_date", "Ngày PO"),
            Column("amount_before_tax", "Tiền trước thuế"),
            Column("vat_amount", "Tiền thuế"),
            Column("total_amount", "Tổng tiền", True),
        ),
        "Dùng để đối chiếu hoá đơn mua vào. Số PO đã có sẽ bị từ chối.",
    ),
    "open_invoices": ImportKind(
        "Hoá đơn còn công nợ",
        (
            Column("direction", "Loại (MUA/BAN)", True),
            Column("party_tax_code", "MST đối tượng", True),
            Column("series", "Ký hiệu"),
            Column("number", "Số hoá đơn", True),
            Column("issue_date", "Ngày hoá đơn", True),
            Column("due_date", "Hạn thanh toán"),
            Column("amount_before_tax", "Tiền trước thuế"),
            Column("vat_amount", "Tiền thuế"),
            Column("total_amount", "Tổng tiền", True),
            Column("paid_amount", "Đã thanh toán"),
        ),
        "Chi tiết công nợ để tính tuổi nợ. Hoá đơn đã hạch toán trong sổ: không tạo bút toán mới.",
    ),
}


class ImportRejected(Exception):
    def __init__(self, errors: list[dict[str, Any]]) -> None:
        super().__init__(f"{len(errors)} lỗi")
        self.errors = errors


def _fold(text: Any) -> str:
    """Header text with accents, case and punctuation removed, for matching."""
    value = unicodedata.normalize("NFD", str(text or "")).replace("đ", "d").replace("Đ", "D")
    value = "".join(ch for ch in value if unicodedata.category(ch) != "Mn")
    return re.sub(r"[^a-z0-9]", "", value.lower())


def template_workbook(kind: str) -> bytes:
    spec = KINDS[kind]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Dữ liệu"
    sheet.append([column.label for column in spec.columns])
    guide = workbook.create_sheet("Hướng dẫn")
    guide.append([spec.title])
    guide.append([spec.guide])
    guide.append(["Cột bắt buộc: " + ", ".join(c.label for c in spec.columns if c.required)])
    guide.append(["Số tiền ghi dạng 1234567 hoặc 1.234.567; ngày dạng dd/mm/yyyy hoặc yyyy-mm-dd."])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def read_rows(kind: str, data: bytes) -> list[tuple[int, dict[str, Any]]]:
    """(Excel row number, {column key: cell}) for every non-empty row of the first sheet."""
    spec = KINDS[kind]
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises several unrelated types for a bad file
        raise ImportRejected([{"row": None, "column": None, "message": f"Không đọc được file Excel: {exc}"}]) from exc
    sheet = workbook.worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    header = next(rows, None) or ()
    wanted = {}
    for column in spec.columns:
        wanted[_fold(column.label)] = column.key
        wanted[_fold(column.key)] = column.key
    positions: dict[str, int] = {}
    for index, cell in enumerate(header):
        key = wanted.get(_fold(cell))
        if key and key not in positions:
            positions[key] = index
    missing = [c.label for c in spec.columns if c.required and c.key not in positions]
    if missing:
        raise ImportRejected([
            {"row": 1, "column": label, "message": "Thiếu cột bắt buộc"} for label in missing
        ])
    result: list[tuple[int, dict[str, Any]]] = []
    for number, values in enumerate(rows, start=2):
        record = {key: values[index] if index < len(values) else None for key, index in positions.items()}
        if all(value in (None, "") for value in record.values()):
            continue
        result.append((number, record))
        if len(result) > MAX_ROWS:
            raise ImportRejected([{"row": number, "column": None, "message": f"Quá {MAX_ROWS} dòng"}])
    return result


class _Collector:
    def __init__(self) -> None:
        self.errors: list[dict[str, Any]] = []

    def check(self, row: int, column: str, parse: Callable[[], Any]) -> Any:
        try:
            return parse()
        except (ValueError, ArithmeticError) as exc:
            self.errors.append({"row": row, "column": column, "message": str(exc)})
            return None

    def add(self, row: int | None, column: str | None, message: str) -> None:
        self.errors.append({"row": row, "column": column, "message": message})


def _text(value: Any, limit: int) -> str:
    text = "" if value is None else str(value).strip()
    if isinstance(value, float) and value.is_integer():
        text = str(int(value))
    return text[:limit]


def _required(value: Any, label: str) -> Any:
    if value in (None, ""):
        raise ValueError(f"{label} không được để trống")
    return value


def _account(value: Any) -> str:
    code = _text(value, 20)
    if not re.fullmatch(r"\d{3,10}", code):
        raise ValueError(f"Số TK không hợp lệ: {value!r}")
    return code


_PARTY_KINDS = {"ncc": "VENDOR", "nhacungcap": "VENDOR", "vendor": "VENDOR",
                "kh": "CUSTOMER", "khachhang": "CUSTOMER", "customer": "CUSTOMER",
                "cahai": "BOTH", "both": "BOTH"}
_DIRECTIONS = {"mua": "IN", "muavao": "IN", "in": "IN", "ban": "OUT", "banra": "OUT", "out": "OUT"}


def _choice(value: Any, choices: dict[str, str], label: str) -> str:
    folded = _fold(value)
    if folded not in choices:
        raise ValueError(f"{label} không hợp lệ: {value!r}")
    return choices[folded]


def import_books(
    db: Session, actor: User, kind: str, filename: str, data: bytes
) -> FinImportBatch:
    """Check the whole file, then write it under one batch. Raises ImportRejected."""
    if kind not in KINDS:
        raise ImportRejected([{"row": None, "column": None, "message": f"Loại dữ liệu không hỗ trợ: {kind}"}])
    rows = read_rows(kind, data)
    if not rows:
        raise ImportRejected([{"row": None, "column": None, "message": "File không có dòng dữ liệu nào"}])
    tenant_id = actor.tenant_id
    accounts = {
        code for (code,) in db.query(FinAccount.code).filter(
            FinAccount.tenant_id == tenant_id, FinAccount.is_active.is_(True)
        )
    }
    parties = {
        tax_code: party_id for party_id, tax_code in db.query(FinParty.id, FinParty.tax_code).filter(
            FinParty.tenant_id == tenant_id, FinParty.tax_code.isnot(None)
        )
    }
    collector = _Collector()
    batch = FinImportBatch(
        id=uuid.uuid4(), tenant_id=tenant_id, kind=kind, filename=filename[:255],
        row_count=len(rows), created_by_id=actor.id,
    )
    builders = {
        "accounts": _accounts,
        "parties": _parties,
        "opening_balances": _ledger_like,
        "ledger": _ledger_like,
        "budgets": _budgets,
        "purchase_orders": _purchase_orders,
        "open_invoices": _open_invoices,
    }
    context = _Context(db, tenant_id, kind, batch.id, accounts, parties, collector)
    pending = builders[kind](context, rows)
    if collector.errors:
        raise ImportRejected(collector.errors[:200])
    db.add(batch)
    db.flush()
    # Changes to rows that already exist are staged as callables so that checking a file
    # never touches the session: a rejected import leaves nothing behind to flush.
    # They run, and are flushed, before the new rows go in: a replaced budget line is
    # deleted first, or its successor would collide with it on the unique key.
    for item in pending:
        if callable(item):
            item()
    db.flush()
    db.add_all(item for item in pending if not callable(item))
    db.flush()
    return batch


@dataclass
class _Context:
    db: Session
    tenant_id: uuid.UUID
    kind: str
    batch_id: uuid.UUID
    accounts: set[str]
    parties: dict[str, uuid.UUID]
    collector: _Collector

    def known_account(self, row: int, value: Any) -> str | None:
        code = self.collector.check(row, "Số TK", lambda: _account(value))
        if code and code not in self.accounts:
            self.collector.add(row, "Số TK", f"TK {code} chưa có trong hệ thống tài khoản")
            return None
        return code

    def party(self, row: int, value: Any, *, required: bool = False) -> uuid.UUID | None:
        tax_code = self.collector.check(row, "MST", lambda: normalize_tax_code(value))
        if tax_code is None:
            if required and value in (None, ""):
                self.collector.add(row, "MST", "MST đối tượng không được để trống")
            return None
        party_id = self.parties.get(tax_code)
        if party_id is None:
            self.collector.add(row, "MST", f"Chưa có đối tượng MST {tax_code}; hãy nhập danh sách đối tượng trước")
        return party_id


def _accounts(ctx: _Context, rows: list[tuple[int, dict[str, Any]]]) -> list[Any]:
    existing = {
        account.code: account
        for account in ctx.db.query(FinAccount).filter(FinAccount.tenant_id == ctx.tenant_id)
    }
    seen: set[str] = set()
    pending: list[Any] = []
    for row, record in rows:
        code = ctx.collector.check(row, "Số TK", lambda: _account(record.get("code")))
        name = ctx.collector.check(row, "Tên TK", lambda: _text(_required(record.get("name"), "Tên TK"), 255))
        if not code or not name:
            continue
        if code in seen:
            ctx.collector.add(row, "Số TK", f"TK {code} lặp lại trong file")
            continue
        seen.add(code)
        if code in existing:
            pending.append(lambda account=existing[code], name=name: (
                setattr(account, "name", name), setattr(account, "is_active", True)
            ))
        else:
            pending.append(FinAccount(
                tenant_id=ctx.tenant_id, code=code, name=name,
                parent_code=parent_code(code), normal_balance=normal_balance(code), is_active=True,
            ))
    return pending


def _parties(ctx: _Context, rows: list[tuple[int, dict[str, Any]]]) -> list[Any]:
    existing = {
        party.tax_code: party
        for party in ctx.db.query(FinParty).filter(
            FinParty.tenant_id == ctx.tenant_id, FinParty.tax_code.isnot(None)
        )
    }
    seen: set[str] = set()
    pending: list[Any] = []
    for row, record in rows:
        kind = ctx.collector.check(row, "Loại", lambda: _choice(record.get("kind"), _PARTY_KINDS, "Loại"))
        tax_code = ctx.collector.check(row, "MST", lambda: normalize_tax_code(record.get("tax_code")))
        name = ctx.collector.check(row, "Tên", lambda: _text(_required(record.get("name"), "Tên"), 255))
        terms = ctx.collector.check(
            row, "Số ngày được nợ",
            lambda: int(record.get("payment_terms_days") or 30),
        )
        if not kind or not name or terms is None:
            continue
        if tax_code:
            if tax_code in seen:
                ctx.collector.add(row, "MST", f"MST {tax_code} lặp lại trong file")
                continue
            seen.add(tax_code)
        values = {
            "kind": kind,
            "name": name,
            "address": _text(record.get("address"), 2000) or None,
            "email": _text(record.get("email"), 255) or None,
            "bank_account": _text(record.get("bank_account"), 50) or None,
            "bank_name": _text(record.get("bank_name"), 255) or None,
            "payment_terms_days": terms,
        }
        party = existing.get(tax_code) if tax_code else None
        if party is not None:
            pending.append(lambda party=party, values=values: [
                setattr(party, field, value) for field, value in values.items() if value is not None
            ])
        else:
            pending.append(FinParty(
                tenant_id=ctx.tenant_id, tax_code=tax_code, is_new=False,
                import_batch_id=ctx.batch_id, **values,
            ))
    return pending


def _amounts(ctx: _Context, row: int, record: dict[str, Any]) -> tuple[Decimal, Decimal] | None:
    debit = ctx.collector.check(row, "Nợ", lambda: to_decimal(record.get("debit")))
    credit = ctx.collector.check(row, "Có", lambda: to_decimal(record.get("credit")))
    if debit is None or credit is None:
        return None
    if debit < 0 or credit < 0:
        ctx.collector.add(row, "Nợ/Có", "Số tiền không được âm")
        return None
    if debit and credit:
        ctx.collector.add(row, "Nợ/Có", "Một dòng chỉ ghi một vế: Nợ hoặc Có")
        return None
    if not debit and not credit:
        ctx.collector.add(row, "Nợ/Có", "Dòng không có số tiền")
        return None
    return debit, credit


_ALREADY_IMPORTED = (
    "File này (hoặc một phần của nó) đã được nhập trước đó. Bỏ các dòng đã có khỏi file, "
    "hoặc hoàn tác lần nhập cũ ở mục Lịch sử nhập rồi nhập lại."
)
# Ledger lines checked against the books per query: an IN list stays a reasonable size.
_CHUNK = 1000


def _reject_known_lines(ctx: _Context, staged: list[tuple[int, FinLedgerLine]]) -> None:
    """Refuse lines the books already hold, before anything is written.

    Every ledger import used to be committed, so loading the same export twice --
    re-sent by email, picked again from Downloads -- counted every voucher twice and no
    balance was right any more. A voucher number on the same day is the same voucher; a
    line without one is the same line when everything it says is the same.
    """
    vouchers: dict[tuple[str, date], int] = {}
    loose: list[tuple[int, FinLedgerLine]] = []
    for row, line in staged:
        if line.voucher_no:
            vouchers.setdefault((line.voucher_no, line.entry_date), row)
        else:
            loose.append((row, line))
    known: set[tuple[str, date]] = set()
    numbers = sorted({number for number, _ in vouchers})
    for index in range(0, len(numbers), _CHUNK):
        known.update(
            (number, day) for number, day in ctx.db.query(FinLedgerLine.voucher_no, FinLedgerLine.entry_date).filter(
                FinLedgerLine.tenant_id == ctx.tenant_id,
                FinLedgerLine.source != "OPENING",
                FinLedgerLine.voucher_no.in_(numbers[index:index + _CHUNK]),
            ).distinct()
        )
    found = False
    for key in sorted(known & vouchers.keys(), key=lambda item: vouchers[item]):
        found = True
        ctx.collector.add(
            vouchers[key], "Số chứng từ", f"Chứng từ {key[0]} ngày {key[1]:%d/%m/%Y} đã có trong sổ"
        )
    if loose:
        days = [line.entry_date for _, line in loose]

        def signature(day, account, debit, credit, party_id, description):
            return day, account, Decimal(debit), Decimal(credit), party_id, description or ""

        existing = {
            signature(*values) for values in ctx.db.query(
                FinLedgerLine.entry_date, FinLedgerLine.account_code, FinLedgerLine.debit,
                FinLedgerLine.credit, FinLedgerLine.party_id, FinLedgerLine.description,
            ).filter(
                FinLedgerLine.tenant_id == ctx.tenant_id,
                FinLedgerLine.source != "OPENING",
                FinLedgerLine.voucher_no.is_(None),
                FinLedgerLine.entry_date >= min(days),
                FinLedgerLine.entry_date <= max(days),
            )
        }
        for row, line in loose:
            if signature(line.entry_date, line.account_code, line.debit, line.credit, line.party_id, line.description) in existing:
                found = True
                ctx.collector.add(row, None, "Dòng này giống hệt một dòng đã có trong sổ")
    if found:
        ctx.collector.add(None, None, _ALREADY_IMPORTED)


def _reject_second_opening(ctx: _Context) -> None:
    """One set of opening balances: every later balance is carried forward from it.

    A second set, of the same period or another, is added on top of the first, and every
    account then opens with both.
    """
    loaded = ctx.db.query(FinLedgerLine.period).filter(
        FinLedgerLine.tenant_id == ctx.tenant_id, FinLedgerLine.source == "OPENING",
    ).order_by(FinLedgerLine.period).first()
    if loaded is not None:
        period = loaded[0]
        ctx.collector.add(
            None, "Kỳ",
            f"Sổ đã có số dư đầu kỳ cho kỳ {period[5:7]}/{period[:4]}; số dư các kỳ sau được tính "
            "tiếp từ đó. Muốn nhập lại thì hoàn tác lần nhập số dư cũ ở mục Lịch sử nhập trước.",
        )


def _ledger_like(ctx: _Context, rows: list[tuple[int, dict[str, Any]]]) -> list[Any]:
    opening = ctx.kind == "opening_balances"
    pending: list[Any] = []
    staged: list[tuple[int, FinLedgerLine]] = []
    totals = [ZERO, ZERO]
    vouchers: dict[str, list[Decimal]] = defaultdict(lambda: [ZERO, ZERO])
    for row, record in rows:
        account = ctx.known_account(row, record.get("account_code"))
        amounts = _amounts(ctx, row, record)
        party_id = ctx.party(row, record.get("party_tax_code")) if record.get("party_tax_code") else None
        if opening:
            period = ctx.collector.check(row, "Kỳ", lambda: normalize_period(record.get("period")))
            entry_date = date(int(period[:4]), int(period[5:]), 1) if period else None
        else:
            entry_date = ctx.collector.check(
                row, "Ngày hạch toán",
                lambda: to_date(_required(record.get("entry_date"), "Ngày hạch toán")),
            )
            period = period_of(entry_date) if entry_date else None
        if not account or not amounts or not period or entry_date is None:
            continue
        debit, credit = amounts
        totals[0] += debit
        totals[1] += credit
        voucher_no = _text(record.get("voucher_no"), 50) or None
        if voucher_no:
            vouchers[voucher_no][0] += debit
            vouchers[voucher_no][1] += credit
        line = FinLedgerLine(
            tenant_id=ctx.tenant_id, period=period, entry_date=entry_date,
            account_code=account, debit=debit, credit=credit, party_id=party_id,
            department=_text(record.get("department"), 50) or None,
            voucher_no=voucher_no,
            description=_text(record.get("description"), 2000) or ("Số dư đầu kỳ" if opening else ""),
            source="OPENING" if opening else "IMPORT",
            import_batch_id=ctx.batch_id,
        )
        pending.append(line)
        staged.append((row, line))
    for voucher_no, (debit, credit) in vouchers.items():
        if debit != credit:
            ctx.collector.add(None, "Số chứng từ", f"Chứng từ {voucher_no} không cân: Nợ {debit} ≠ Có {credit}")
    if not ctx.collector.errors and totals[0] != totals[1]:
        ctx.collector.add(None, None, f"File không cân: tổng Nợ {totals[0]} ≠ tổng Có {totals[1]}")
    if opening:
        _reject_second_opening(ctx)
    elif staged:
        _reject_known_lines(ctx, staged)
    return pending


def _budgets(ctx: _Context, rows: list[tuple[int, dict[str, Any]]]) -> list[Any]:
    existing = {
        (budget.department, budget.account_code, budget.period): budget
        for budget in ctx.db.query(FinBudget).filter(FinBudget.tenant_id == ctx.tenant_id)
    }
    pending: list[Any] = []
    seen: set[tuple[str, str, str]] = set()
    for row, record in rows:
        department = ctx.collector.check(
            row, "Phòng ban", lambda: _text(_required(record.get("department"), "Phòng ban"), 50).upper()
        )
        account = ctx.known_account(row, record.get("account_code"))
        period = ctx.collector.check(row, "Kỳ", lambda: normalize_period(record.get("period")))
        amount = ctx.collector.check(row, "Số tiền", lambda: to_decimal(record.get("amount")))
        if not department or not account or not period or amount is None:
            continue
        key = (department, account, period)
        if key in seen:
            ctx.collector.add(row, "Phòng ban", "Dòng ngân sách lặp lại trong file")
            continue
        seen.add(key)
        if key in existing:
            # Replacing a line moves it to this batch, so undoing the batch removes it.
            pending.append(lambda old=existing[key]: ctx.db.delete(old))
        pending.append(FinBudget(
            tenant_id=ctx.tenant_id, department=department, account_code=account,
            period=period, amount=amount, import_batch_id=ctx.batch_id,
        ))
    return pending


def _invoice_amounts(ctx: _Context, row: int, record: dict[str, Any]) -> tuple[Decimal, Decimal, Decimal] | None:
    total = ctx.collector.check(row, "Tổng tiền", lambda: to_decimal(_required(record.get("total_amount"), "Tổng tiền")))
    before = ctx.collector.check(row, "Tiền trước thuế", lambda: to_decimal(record.get("amount_before_tax")))
    vat = ctx.collector.check(row, "Tiền thuế", lambda: to_decimal(record.get("vat_amount")))
    if total is None or before is None or vat is None:
        return None
    if not before and not vat:
        before = total
    if before + vat != total:
        ctx.collector.add(row, "Tổng tiền", f"Tiền trước thuế + thuế ({before + vat}) ≠ tổng tiền ({total})")
        return None
    return before, vat, total


def _purchase_orders(ctx: _Context, rows: list[tuple[int, dict[str, Any]]]) -> list[Any]:
    existing = {
        number for (number,) in ctx.db.query(FinPurchaseOrder.po_number).filter(
            FinPurchaseOrder.tenant_id == ctx.tenant_id
        )
    }
    pending: list[Any] = []
    for row, record in rows:
        number = ctx.collector.check(row, "Số PO", lambda: _text(_required(record.get("po_number"), "Số PO"), 50))
        party_id = ctx.party(row, record.get("party_tax_code"), required=True)
        order_date = ctx.collector.check(row, "Ngày PO", lambda: to_date(record.get("order_date")))
        amounts = _invoice_amounts(ctx, row, record)
        if not number or not party_id or not amounts:
            continue
        if number in existing:
            ctx.collector.add(row, "Số PO", f"PO {number} đã có")
            continue
        existing.add(number)
        before, vat, total = amounts
        pending.append(FinPurchaseOrder(
            tenant_id=ctx.tenant_id, po_number=number, party_id=party_id, order_date=order_date,
            amount_before_tax=before, vat_amount=vat, total_amount=total,
            import_batch_id=ctx.batch_id,
        ))
    return pending


def _open_invoices(ctx: _Context, rows: list[tuple[int, dict[str, Any]]]) -> list[Any]:
    parties = {
        party.id: party for party in ctx.db.query(FinParty).filter(FinParty.tenant_id == ctx.tenant_id)
    }
    pending: list[Any] = []
    seen: set[tuple[str, str, str, str]] = set()
    for row, record in rows:
        direction = ctx.collector.check(row, "Loại", lambda: _choice(record.get("direction"), _DIRECTIONS, "Loại"))
        party_id = ctx.party(row, record.get("party_tax_code"), required=True)
        number = ctx.collector.check(row, "Số hoá đơn", lambda: _text(_required(record.get("number"), "Số hoá đơn"), 20))
        issue_date = ctx.collector.check(row, "Ngày hoá đơn", lambda: to_date(_required(record.get("issue_date"), "Ngày hoá đơn")))
        due_date = ctx.collector.check(row, "Hạn thanh toán", lambda: to_date(record.get("due_date")))
        paid = ctx.collector.check(row, "Đã thanh toán", lambda: to_decimal(record.get("paid_amount")))
        amounts = _invoice_amounts(ctx, row, record)
        if not direction or not party_id or not number or not issue_date or not amounts or paid is None:
            continue
        before, vat, total = amounts
        if paid < 0 or paid > total:
            ctx.collector.add(row, "Đã thanh toán", "Số đã thanh toán phải từ 0 đến tổng tiền")
            continue
        party = parties[party_id]
        series = _text(record.get("series"), 20)
        seller_tax_code = party.tax_code if direction == "IN" else ""
        key = (direction, seller_tax_code or "", series, number)
        if key in seen:
            ctx.collector.add(row, "Số hoá đơn", "Hoá đơn lặp lại trong file")
            continue
        seen.add(key)
        clash = ctx.db.query(FinInvoice.id).filter(
            FinInvoice.tenant_id == ctx.tenant_id,
            FinInvoice.direction == direction,
            FinInvoice.seller_tax_code == (seller_tax_code or ""),
            FinInvoice.series == series,
            FinInvoice.number == number,
        ).first()
        if clash:
            ctx.collector.add(row, "Số hoá đơn", f"Hoá đơn {series} {number} đã có")
            continue
        invoice_id = uuid.uuid4()
        pending.append(FinInvoice(
            id=invoice_id, tenant_id=ctx.tenant_id, direction=direction, party_id=party_id,
            seller_tax_code=seller_tax_code or "", seller_name=party.name if direction == "IN" else "",
            buyer_tax_code=party.tax_code if direction == "OUT" else None,
            buyer_name=party.name if direction == "OUT" else None,
            series=series, number=number, issue_date=issue_date,
            due_date=due_date or date.fromordinal(issue_date.toordinal() + party.payment_terms_days),
            amount_before_tax=before, vat_amount=vat, total_amount=total,
            # Already in the books: the ledger import carries its posting.
            status="PAID" if paid == total else "POSTED",
            source_format="IMPORT", import_batch_id=ctx.batch_id,
        ))
        if paid:
            pending.append(FinPayment(
                tenant_id=ctx.tenant_id, invoice_id=invoice_id, party_id=party_id,
                direction="PAY" if direction == "IN" else "RECEIVE", amount=paid,
                status="PAID", reference="Số đã thanh toán khi nhập", import_batch_id=ctx.batch_id,
            ))
    return pending


def undo_batch(db: Session, batch: FinImportBatch) -> int:
    """Remove every row an undoable batch wrote. Returns how many rows went."""
    if not KINDS[batch.kind].undoable:
        raise ValueError("Không thể hoàn tác lần nhập danh mục (tài khoản, đối tượng): dữ liệu đã được cập nhật đè")
    invoice_ids = [
        invoice_id for (invoice_id,) in db.query(FinInvoice.id).filter(FinInvoice.import_batch_id == batch.id)
    ]
    if invoice_ids:
        # Payments cascade with their invoice: one recorded after the import (a voucher,
        # a receipt) would silently disappear with it.
        later_payment = db.query(FinPayment.id).filter(
            FinPayment.invoice_id.in_(invoice_ids),
            (FinPayment.import_batch_id.is_(None)) | (FinPayment.import_batch_id != batch.id),
        ).first()
        drafted = db.query(FinJournalEntry.id).filter(FinJournalEntry.invoice_id.in_(invoice_ids)).first()
        if later_payment or drafted:
            raise ValueError("Không thể hoàn tác: đã có thanh toán hoặc bút toán gắn với hoá đơn của lần nhập này")
    removed = 0
    for model in (FinPayment, FinLedgerLine, FinBudget):
        removed += db.query(model).filter(model.import_batch_id == batch.id).delete(synchronize_session=False)
    for model in (FinInvoice, FinPurchaseOrder):
        removed += db.query(model).filter(model.import_batch_id == batch.id).delete(synchronize_session=False)
    batch.status = "ROLLED_BACK"
    db.flush()
    return removed
