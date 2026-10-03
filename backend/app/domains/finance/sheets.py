"""Spreadsheets someone uploads to analyse: read once, then computed on with fixed operations.

A real cost sheet is rarely a clean table: a title and a few blank lines above the column
titles, "Tổng cộng" rows between the groups, amounts typed as text. Reading it is the part
that goes wrong silently -- a total row counted twice is a figure that looks right -- so
the reading is shown back to the uploader (header row, column kinds, the rows left out)
and they can correct it before anything is computed.

The computing is deterministic: a sum, count, average, minimum or maximum of one column,
optionally grouped by another (or by month, quarter, year of a date column) and filtered.
The Finance agent picks the operation and the columns; it never adds anything up itself.
"""

from __future__ import annotations

import csv
import io
import re
import unicodedata
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.domains.finance.money import CENT, ZERO, AmountError, plain, to_date, to_decimal
from app.domains.finance.storage import read_stored_file, save_sheet_file

MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 20_000
MAX_COLUMNS = 60
HEADER_SCAN = 30
MAX_GROUPS = 50
KINDS = ("NUMBER", "DATE", "TEXT")
OPERATIONS = ("sum", "count", "average", "min", "max")
PERIODS = ("day", "month", "quarter", "year")
# A column is of a kind when this share of its filled cells reads as that kind.
KIND_SHARE = 0.8
EMPTY_GROUP = "(trống)"


class SheetRejected(ValueError):
    """The file or the request cannot be read; the message says why, in Vietnamese."""


def _fold(value: Any) -> str:
    text = unicodedata.normalize("NFD", str(value or "")).replace("đ", "d").replace("Đ", "D")
    text = "".join(char for char in text if unicodedata.category(char) != "Mn")
    return re.sub(r"\s+", " ", text).strip().lower()


# "Tổng cộng", "Tổng chi phí quý 1", "Cộng", "Cộng quý 3", "Total". Only as the first filled
# cell of a row with no date. "Cộng" folds like "Công", so it counts only alone or before a
# period word: "Công cụ dụng cụ", "Công tác phí" are costs, not totals.
_TOTAL_LABEL = re.compile(
    r"^(tong|(sub ?|grand )?total)\b"
    r"|^cong\s*(:|$|(quy|thang|nam|ky|phat sinh|luy ke|chung|nhom)\b)"
)


def _total_label(values: list[Any]) -> str | None:
    filled = [value for value in values if _filled(value)]
    if not filled or any(_as_date(value) is not None for value in filled):
        return None
    first = filled[0]
    return str(first).strip() if _as_number(first) is None and _TOTAL_LABEL.match(_fold(first)) else None


# ------------------------------------------------------------------------------ reading


def read_grids(data: bytes, filename: str) -> dict[str, list[list[Any]]]:
    """Every sheet of the file as rows of cell values."""
    if len(data) > MAX_BYTES:
        raise SheetRejected("File lớn hơn 10 MB")
    name = filename.lower()
    if name.endswith(".csv"):
        return {"CSV": _csv_grid(data)}
    if name.endswith(".xls"):
        raise SheetRejected("File .xls (Excel 97-2003) chưa đọc được; hãy lưu lại dưới dạng .xlsx")
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises several unrelated types for a bad file
        raise SheetRejected(f"Không đọc được file Excel: {exc}") from exc
    grids = {}
    for sheet in workbook.worksheets:
        grid = []
        for values in sheet.iter_rows(values_only=True):
            grid.append(list(values[:MAX_COLUMNS]))
            if len(grid) > MAX_ROWS + HEADER_SCAN:
                raise SheetRejected(f"Sheet {sheet.title} có hơn {MAX_ROWS} dòng")
        grids[sheet.title] = grid
    workbook.close()
    if not grids:
        raise SheetRejected("File không có sheet nào")
    return grids


def _csv_grid(data: bytes) -> list[list[Any]]:
    for encoding in ("utf-8-sig", "cp1258", "cp1252"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise SheetRejected("Không đọc được bảng mã của file CSV")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    grid = [row[:MAX_COLUMNS] for row in csv.reader(io.StringIO(text), dialect)]
    if len(grid) > MAX_ROWS + HEADER_SCAN:
        raise SheetRejected(f"File có hơn {MAX_ROWS} dòng")
    return grid


def _filled(value: Any) -> bool:
    return value is not None and str(value).strip() != ""


def _as_number(value: Any) -> Decimal | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float, Decimal)):
        return to_decimal(value)
    text = str(value).strip().replace("₫", "")
    if not re.search(r"\d", text) or re.search(r"[A-Za-zÀ-ỹ]", re.sub(r"(?i)vn[dđ]|đ", "", text)) or "/" in text:
        return None
    try:
        return to_decimal(text)
    except AmountError:
        return None


def _as_date(value: Any) -> date | None:
    if isinstance(value, (datetime, date)):
        return value.date() if isinstance(value, datetime) else value
    text = str(value).strip()
    if not re.match(r"^\d{1,4}[/-]\d{1,2}[/-]\d{1,4}", text):
        return None
    try:
        return to_date(text)
    except ValueError:
        return None


def detect_header(grid: list[list[Any]]) -> int:
    """0-based index of the likeliest column-title row among the first rows."""
    best, best_score = 0, -1.0
    for index, row in enumerate(grid[:HEADER_SCAN]):
        titles = sum(1 for value in row if _filled(value) and _as_number(value) is None and _as_date(value) is None)
        if titles < 2:
            continue
        below = grid[index + 1:index + 4]
        numbers_below = any(_as_number(value) is not None for line in below for value in line if _filled(value))
        score = titles + (0.5 if numbers_below else 0)
        if score > best_score:
            best, best_score = index, score
    return best


def _kind_of(values: list[Any]) -> str:
    filled = [value for value in values if _filled(value)]
    if not filled:
        return "TEXT"
    dates = sum(1 for value in filled if _as_date(value) is not None)
    if dates >= KIND_SHARE * len(filled):
        return "DATE"
    numbers = sum(1 for value in filled if _as_number(value) is not None)
    return "NUMBER" if numbers >= KIND_SHARE * len(filled) else "TEXT"


def _cell(value: Any, kind: str) -> Any:
    if not _filled(value):
        return None
    if kind == "NUMBER":
        number = _as_number(value)
        return plain(number) if number is not None else None
    if kind == "DATE":
        day = _as_date(value)
        return day.isoformat() if day else None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def parse_grid(
    grid: list[list[Any]],
    *,
    header_row: int | None = None,
    kinds: dict[int, str] | None = None,
    skip_totals: bool = True,
) -> dict[str, Any]:
    """Columns, rows and the rows left out, from one sheet's cells.

    `header_row` is 1-based like Excel's own numbering; None lets the reader find it.
    `kinds` overrides the detected kind of a column, by its index.
    """
    if not grid:
        raise SheetRejected("Sheet trống")
    header_index = detect_header(grid) if header_row is None else header_row - 1
    if not 0 <= header_index < len(grid):
        raise SheetRejected(f"Dòng tiêu đề {header_index + 1} nằm ngoài sheet")
    header = grid[header_index]
    body = grid[header_index + 1:]
    if len(body) > MAX_ROWS:
        raise SheetRejected(f"Sheet có hơn {MAX_ROWS} dòng")
    width = max([len(header)] + [len(row) for row in body[:200]])
    used = [
        index for index in range(min(width, MAX_COLUMNS))
        if (index < len(header) and _filled(header[index]))
        or any(index < len(row) and _filled(row[index]) for row in body)
    ]
    if not used:
        raise SheetRejected("Không tìm thấy cột dữ liệu nào dưới dòng tiêu đề")
    names: list[str] = []
    for index in used:
        title = str(header[index]).strip() if index < len(header) and _filled(header[index]) else f"Cột {get_column_letter(index + 1)}"
        name, copy = title, 2
        while _fold(name) in {_fold(existing) for existing in names}:
            name, copy = f"{title} ({copy})", copy + 1
        names.append(name)

    kept: list[tuple[int, list[Any]]] = []
    skipped: list[dict[str, Any]] = []
    for offset, row in enumerate(body):
        number = header_index + 2 + offset
        values = [row[index] if index < len(row) else None for index in used]
        if not any(_filled(value) for value in values):
            continue
        label = _total_label(values) if skip_totals else None
        if label is not None:
            skipped.append({"row": number, "reason": "TOTAL", "label": label[:80]})
            continue
        kept.append((number, values))

    detected = [_kind_of([values[position] for _, values in kept]) for position in range(len(used))]
    columns = []
    for position, index in enumerate(used):
        kind = (kinds or {}).get(index) or detected[position]
        if kind not in KINDS:
            raise SheetRejected(f"Kiểu cột không hợp lệ: {kind}")
        invalid = sum(
            1 for _, values in kept
            if _filled(values[position]) and _cell(values[position], kind) is None
        )
        columns.append({"index": index, "name": names[position], "kind": kind, "detected": detected[position], "invalid": invalid})
    rows = [
        [number] + [_cell(values[position], columns[position]["kind"]) for position in range(len(used))]
        for number, values in kept
    ]
    return {"header_row": header_index + 1, "columns": columns, "rows": rows, "skipped_rows": skipped}


# ------------------------------------------------------------------------------ analysing


def resolve_column(columns: list[dict[str, Any]], name: str | None, *, purpose: str) -> tuple[int, dict[str, Any]]:
    """(position in a stored row, column) for a column named by the user or the model."""
    wanted = _fold(name)
    exact = [item for item in enumerate(columns) if _fold(item[1]["name"]) == wanted]
    partial = [item for item in enumerate(columns) if wanted and wanted in _fold(item[1]["name"])]
    found = exact or (partial if len(partial) == 1 else [])
    if not found:
        available = ", ".join(f"“{column['name']}”" for column in columns)
        raise SheetRejected(f"Không có cột “{name}” để {purpose}. Các cột: {available}")
    position, column = found[0]
    return position + 1, column  # stored rows start with the Excel row number


def _matches(value: Any, kind: str, op: str, wanted: str) -> bool:
    if value is None:
        return False
    if kind == "NUMBER":
        left, right = Decimal(value), _as_number(wanted)
        if right is None:
            raise SheetRejected(f"“{wanted}” không phải là số")
    elif kind == "DATE":
        left, right = date.fromisoformat(value), _as_date(wanted)
        if right is None:
            raise SheetRejected(f"“{wanted}” không phải là ngày (dd/mm/yyyy)")
    else:
        left, right = _fold(value), _fold(wanted)
        if op == "contains":
            return right in left
    if op == "eq":
        return left == right
    if op == "gte":
        return left >= right
    if op == "lte":
        return left <= right
    raise SheetRejected(f"Phép lọc “{op}” chỉ dùng được cho cột chữ")


def _group_key(value: Any, kind: str, period: str | None) -> str:
    if value is None:
        return EMPTY_GROUP
    if kind != "DATE":
        return str(value)
    day = date.fromisoformat(value)
    if period == "year":
        return f"{day.year}"
    if period == "quarter":
        return f"{day.year}-Q{(day.month - 1) // 3 + 1}"
    if period == "day":
        return day.isoformat()
    return f"{day.year:04d}-{day.month:02d}"


def _reduce(operation: str, values: list[Decimal], count: int) -> Decimal | int | None:
    if operation == "count":
        return count
    if not values:
        return None
    if operation == "sum":
        return sum(values, ZERO)
    if operation == "average":
        return (sum(values, ZERO) / len(values)).quantize(CENT)
    return min(values) if operation == "min" else max(values)


def _out(value: Decimal | int | None) -> str | int | None:
    return plain(value) if isinstance(value, Decimal) else value


def analyze(
    sheet: Any,
    *,
    operation: str,
    value_column: str | None = None,
    group_by: str | None = None,
    period: str | None = None,
    filters: list[dict[str, Any]] | None = None,
) -> dict[str, Any]:
    """One fixed computation over the stored rows of `sheet` (a FinSheet)."""
    if operation not in OPERATIONS:
        raise SheetRejected(f"Phép tính “{operation}” không hỗ trợ; dùng: {', '.join(OPERATIONS)}")
    columns = sheet.columns or []
    value_position = value = None
    if value_column or operation != "count":
        value_position, value = resolve_column(columns, value_column, purpose="tính")
        if operation != "count" and value["kind"] != "NUMBER":
            raise SheetRejected(f"Cột “{value['name']}” không phải cột số nên không tính {operation} được")
    group_position = group = None
    if group_by:
        group_position, group = resolve_column(columns, group_by, purpose="nhóm")
    if period and (group is None or group["kind"] != "DATE"):
        raise SheetRejected("Nhóm theo tháng/quý/năm cần một cột ngày ở group_by")
    conditions = []
    for item in filters or []:
        position, column = resolve_column(columns, item.get("column"), purpose="lọc")
        conditions.append((position, column, str(item.get("op") or "eq"), str(item.get("value") or "")))

    numbers: dict[str, list[Decimal]] = defaultdict(list)
    counts: dict[str, int] = defaultdict(int)
    every: list[Decimal] = []
    used = 0
    for row in sheet.rows or []:
        if not all(_matches(row[position], column["kind"], op, wanted) for position, column, op, wanted in conditions):
            continue
        cell = row[value_position] if value_position is not None else None
        if value is not None and cell is None:
            continue
        used += 1
        key = _group_key(row[group_position], group["kind"], period) if group is not None else ""
        counts[key] += 1
        if value is not None and value["kind"] == "NUMBER":
            numbers[key].append(Decimal(cell))
            every.append(Decimal(cell))

    total = _reduce(operation, every, used)
    groups = [
        {"group": key, "value": _reduce(operation, numbers[key], counts[key]), "rows": counts[key]}
        for key in counts
    ] if group is not None else []
    chronological = group is not None and group["kind"] == "DATE"
    groups.sort(key=(lambda item: item["group"]) if chronological else (lambda item: (item["value"] is None, -(item["value"] or 0))))
    folded = 0
    if len(groups) > MAX_GROUPS and operation in {"sum", "count"} and not chronological:
        rest = groups[MAX_GROUPS - 1:]
        folded = len(rest)
        groups = groups[:MAX_GROUPS - 1] + [{
            "group": "Khác",
            "value": sum((item["value"] for item in rest), ZERO if operation == "sum" else 0),
            "rows": sum(item["rows"] for item in rest),
        }]
    shares = operation in {"sum", "count"} and total and all((item["value"] or 0) >= 0 for item in groups)
    return {
        "sheet_id": str(sheet.id),
        "file": sheet.filename,
        "sheet": sheet.sheet_name,
        "operation": operation,
        "value_column": value["name"] if value else None,
        "group_by": group["name"] if group else None,
        "period": period if group is not None and group["kind"] == "DATE" else None,
        "filters": [{"column": column["name"], "op": op, "value": wanted} for _, column, op, wanted in conditions],
        "groups": [
            {
                **item,
                "value": _out(item["value"]),
                # Given outright so a reply never has to work a share out itself.
                "share_percent": plain(Decimal(item["value"]) * 100 / Decimal(total)) if shares and item["value"] is not None else None,
            }
            for item in groups
        ],
        "groups_folded_into_other": folded,
        "total": _out(total),
        "rows_used": used,
        "rows_left_out_as_totals": len(sheet.skipped_rows or []),
        "source": f"file “{sheet.filename}”, sheet {sheet.sheet_name}, {used} dòng",
    }


def describe(sheet: Any) -> dict[str, Any]:
    """What the agent needs to choose an analysis: the columns and how big the sheet is."""
    return {
        "sheet_id": str(sheet.id),
        "file": sheet.filename,
        "sheet": sheet.sheet_name,
        "uploaded_at": sheet.created_at.isoformat() if sheet.created_at else None,
        "columns": [{"name": column["name"], "kind": column["kind"]} for column in sheet.columns or []],
        "row_count": sheet.row_count,
        "rows_left_out_as_totals": len(sheet.skipped_rows or []),
    }


# ------------------------------------------------------------------------------ records


def _apply(sheet: Any, grids: dict[str, list[list[Any]]], *, sheet_name: str | None, header_row: int | None,
           kinds: dict[int, str] | None, skip_totals: bool) -> None:
    name = sheet_name or next((title for title, grid in grids.items() if any(any(_filled(v) for v in row) for row in grid)), next(iter(grids)))
    if name not in grids:
        raise SheetRejected(f"Không có sheet “{name}”")
    parsed = parse_grid(grids[name], header_row=header_row, kinds=kinds, skip_totals=skip_totals)
    sheet.sheet_names = list(grids)
    sheet.sheet_name = name
    sheet.header_row = parsed["header_row"]
    sheet.columns = parsed["columns"]
    sheet.rows = parsed["rows"]
    sheet.skipped_rows = parsed["skipped_rows"]
    sheet.skip_totals = skip_totals
    sheet.row_count = len(parsed["rows"])


def create_sheet(db: Any, actor: Any, filename: str, data: bytes) -> Any:
    from app.models.models import FinSheet

    grids = read_grids(data, filename)
    sheet = FinSheet(tenant_id=actor.tenant_id, created_by_id=actor.id, filename=filename[:255], storage_key="")
    _apply(sheet, grids, sheet_name=None, header_row=None, kinds=None, skip_totals=True)
    db.add(sheet)
    db.flush()
    sheet.storage_key = save_sheet_file(tenant_id=actor.tenant_id, sheet_id=sheet.id, filename=filename, data=data)
    return sheet


def reread_sheet(sheet: Any, *, sheet_name: str | None, header_row: int | None,
                 kinds: dict[int, str] | None, skip_totals: bool) -> Any:
    """Read the stored file again with the uploader's corrections."""
    grids = read_grids(read_stored_file(sheet.storage_key), sheet.filename)
    changed_sheet = sheet_name is not None and sheet_name != sheet.sheet_name
    _apply(
        sheet, grids,
        sheet_name=sheet_name or sheet.sheet_name,
        # Another sheet has its own title row; a correction made for the old one does not carry.
        header_row=None if changed_sheet and header_row is None else (header_row or sheet.header_row),
        kinds=None if changed_sheet else kinds,
        skip_totals=skip_totals,
    )
    return sheet


def own_sheets(db: Any, actor: Any) -> Any:
    from app.models.models import FinSheet

    return db.query(FinSheet).filter(FinSheet.tenant_id == actor.tenant_id, FinSheet.created_by_id == actor.id)


def own_sheet(db: Any, actor: Any, sheet_id: Any) -> Any:
    """The uploader's sheet, or the latest one when no id is given; None otherwise."""
    from app.models.models import FinSheet

    query = own_sheets(db, actor)
    if sheet_id:
        return query.filter(FinSheet.id == sheet_id).first()
    return query.order_by(FinSheet.created_at.desc()).first()


def preview(sheet: Any, limit: int = 30) -> dict[str, Any]:
    return {
        **describe(sheet),
        "sheet_names": sheet.sheet_names,
        "header_row": sheet.header_row,
        "skip_totals": sheet.skip_totals,
        "columns": sheet.columns,
        "rows": (sheet.rows or [])[:limit],
        "skipped_rows": sheet.skipped_rows,
    }
