"""Build the redline report for a reviewed contract.

This produces a *new* document rather than editing the uploaded one. The original
bytes are never stored, `extract_file_text` is one-way (formatting, numbering and
sections are gone by the time the analyzer sees the text), most uploads are not
DOCX at all, and python-docx has no tracked-changes API. A report is what can
honestly be produced -- which is also what the UI promises: "AI không thay đổi
file gốc; mỗi đề xuất cần được xác nhận".

Struck-through original beside underlined replacement reads as a redline to a
lawyer without fabricating `w:ins`/`w:del` revisions against a document we do not
have.
"""

from __future__ import annotations

import io
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from app.services.docx_style import (
    set_cell_shading,
    set_docx_font,
    set_table_geometry,
    style_docx_table,
)


ACCEPTED_DECISIONS = {"ACCEPTED", "EDITED"}
SEVERITY_COLORS = {
    "CRITICAL": "9B1C1C",
    "HIGH": "B54708",
    "MEDIUM": "B58B00",
    "LOW": "175CD3",
}
DECISION_LABELS = {
    "ACCEPTED": "ĐÃ CHẤP NHẬN ĐỀ XUẤT",
    "EDITED": "ĐÃ CHỈNH SỬA ĐỀ XUẤT",
    "REJECTED": "ĐÃ TỪ CHỐI",
}


def _safe_stem(document_name: str) -> str:
    stem = Path(str(document_name or "contract")).stem
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "-", stem).strip("-._")
    return cleaned[:60] or "contract"


def applied_revision(finding: dict[str, Any], decision: dict[str, Any]) -> str:
    """What the contract should say: the reviewer's wording wins over the AI's."""
    if decision.get("decision") == "EDITED" and (decision.get("revised_text") or "").strip():
        return str(decision["revised_text"]).strip()
    return str(finding.get("suggested_revision") or "").strip()


def build_redline_docx(
    review_result: dict[str, Any],
    decisions: list[dict[str, Any]],
    *,
    document_name: str,
    generated_by: str,
    review_id: str,
) -> tuple[bytes, str]:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    by_key = {str(item.get("finding_key")): item for item in decisions}
    findings = list(review_result.get("findings") or [])
    applied = [
        finding for finding in findings
        if by_key.get(str(finding.get("finding_key")), {}).get("decision") in ACCEPTED_DECISIONS
    ]
    outstanding = [finding for finding in findings if finding not in applied]

    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    for margin in ("top_margin", "right_margin", "bottom_margin", "left_margin"):
        setattr(section, margin, Inches(1))

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    for style_name, size, color in (("Heading 1", 13, "2E74B5"), ("Heading 2", 12, "1F4D78")):
        style = document.styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    set_docx_font(
        header.add_run("AI WORKFORCE  |  LEGAL REDLINE REPORT"), size=8, bold=True, color="667085"
    )
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_docx_font(
        footer.add_run("Tài liệu nội bộ · Không thay thế ý kiến pháp lý"), size=8, color="98A2B3"
    )

    title = document.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    set_docx_font(
        title.add_run("BÁO CÁO RÀ SOÁT & ĐỀ XUẤT CHỈNH SỬA"),
        size=20, bold=True, color="0B2545",
    )
    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    set_docx_font(subtitle.add_run(str(document_name)), size=12, color="475467")

    generated_at = datetime.now(timezone.utc).strftime("%d/%m/%Y %H:%M UTC")
    meta_rows = [
        ("Loại hợp đồng", str(review_result.get("contract_type_label") or review_result.get("contract_type") or "—")),
        ("Góc nhìn rà soát", str(review_result.get("represented_party_label") or review_result.get("represented_party") or "—")),
        ("Điểm rủi ro", f"{review_result.get('risk_score', 0)}/100 ({review_result.get('risk_level', 'LOW')})"),
        ("Đề xuất được áp dụng", f"{len(applied)}/{len(findings)}"),
        ("Người rà soát", str(generated_by)),
        ("Thời điểm xuất", generated_at),
        ("Mã bản rà soát", str(review_id)),
    ]
    metadata = document.add_table(rows=0, cols=2)
    for label, value in meta_rows:
        cells = metadata.add_row().cells
        cells[0].text = label
        cells[1].text = value
    style_docx_table(metadata, [2700, 6660], header=False)
    for row in metadata.rows:
        set_cell_shading(row.cells[0], "E8EEF5")
        for run in row.cells[0].paragraphs[0].runs:
            set_docx_font(run, size=9, bold=True, color="1F4D78")

    disclaimer = document.add_paragraph()
    disclaimer.paragraph_format.space_before = Pt(10)
    set_docx_font(
        disclaimer.add_run(str(review_result.get("review_disclaimer") or "")),
        size=9, italic=True, color="667085",
    )

    counts = review_result.get("severity_counts") or {}
    document.add_heading("1. TỔNG QUAN RỦI RO", level=1)
    summary = document.add_table(rows=1, cols=4)
    for index, label in enumerate(("CRITICAL", "HIGH", "MEDIUM", "LOW")):
        summary.rows[0].cells[index].text = label
    values = summary.add_row().cells
    for index, label in enumerate(("CRITICAL", "HIGH", "MEDIUM", "LOW")):
        values[index].text = str(counts.get(label, 0))
    style_docx_table(summary, [2340] * 4)

    extra = document.add_paragraph()
    set_docx_font(
        extra.add_run(
            f"Thiếu điều khoản: {review_result.get('missing_clauses_count', 0)} · "
            f"Mâu thuẫn nội bộ: {review_result.get('internal_conflicts_count', 0)} · "
            f"Vi phạm policy: {review_result.get('policy_violations_count', 0)}"
        ),
        size=10,
    )

    checklist = review_result.get("checklist") or []
    if checklist:
        document.add_heading("2. CHECKLIST THEO LOẠI HỢP ĐỒNG", level=1)
        table = document.add_table(rows=1, cols=3)
        for index, label in enumerate(("Nhóm điều khoản", "Trạng thái", "Mức độ nếu thiếu")):
            table.rows[0].cells[index].text = label
        for item in checklist:
            cells = table.add_row().cells
            cells[0].text = str(item.get("label", ""))
            cells[1].text = "Có" if item.get("status") == "PRESENT" else "Thiếu"
            cells[2].text = "—" if item.get("status") == "PRESENT" else str(item.get("severity_if_missing", ""))
        style_docx_table(table, [4680, 2340, 2340])

    document.add_heading("3. ĐỀ XUẤT CHỈNH SỬA ĐÃ ĐƯỢC CHẤP NHẬN", level=1)
    if not applied:
        set_docx_font(
            document.add_paragraph().add_run("Chưa có đề xuất nào được chấp nhận."),
            size=10, italic=True, color="667085",
        )
    for index, finding in enumerate(applied, 1):
        decision = by_key.get(str(finding.get("finding_key")), {})
        clause = finding.get("clause")
        heading_text = (
            f"3.{index}. {finding.get('issue', '')}"
            if clause == "MISSING"
            else f"3.{index}. Điều {clause} · {finding.get('issue', '')}"
        )
        heading = document.add_heading(heading_text, level=2)
        heading.paragraph_format.keep_with_next = True

        severity_line = document.add_paragraph()
        set_docx_font(
            severity_line.add_run(f"[{finding.get('severity', 'LOW')}] "),
            size=9, bold=True,
            color=SEVERITY_COLORS.get(str(finding.get("severity")), "667085"),
        )
        set_docx_font(
            severity_line.add_run(str(finding.get("finding_type", ""))), size=9, color="667085"
        )

        for label, value in (
            ("Lý do: ", finding.get("reason", "")),
            ("Khuyến nghị: ", finding.get("recommendation", "")),
        ):
            paragraph = document.add_paragraph()
            set_docx_font(paragraph.add_run(label), size=10, bold=True)
            set_docx_font(paragraph.add_run(str(value)), size=10)

        comparison = document.add_table(rows=2, cols=2)
        comparison.rows[0].cells[0].text = "NGUYÊN VĂN"
        comparison.rows[0].cells[1].text = "ĐỀ XUẤT THAY THẾ"
        original_cell = comparison.rows[1].cells[0]
        revised_cell = comparison.rows[1].cells[1]
        original_cell.text = ""
        revised_cell.text = ""
        set_docx_font(
            original_cell.paragraphs[0].add_run(
                str(finding.get("original_text") or "Không có trong hợp đồng gốc.")
            ),
            size=9, color="9B1C1C", strike=True,
        )
        set_docx_font(
            revised_cell.paragraphs[0].add_run(applied_revision(finding, decision)),
            size=9, color="176B45", underline=True,
        )
        style_docx_table(comparison, [4680, 4680])

        stamp = document.add_paragraph()
        set_docx_font(
            stamp.add_run(
                f"{DECISION_LABELS.get(decision.get('decision'), decision.get('decision', ''))}"
                f" · {decision.get('decided_by_name') or 'Không rõ'}"
                f" · {decision.get('updated_at') or ''}"
            ),
            size=8, bold=True, color="176B45",
        )
        if decision.get("comment"):
            set_docx_font(
                document.add_paragraph().add_run(f"Ghi chú: {decision['comment']}"),
                size=9, italic=True, color="475467",
            )
        sources = finding.get("sources") or []
        if sources:
            set_docx_font(
                document.add_paragraph().add_run(
                    "Nguồn: " + "; ".join(str(source.get("title", "")) for source in sources)
                ),
                size=8, color="667085",
            )

    document.add_page_break()
    document.add_heading("PHỤ LỤC A — PHÁT HIỆN CHƯA ÁP DỤNG", level=1)
    if outstanding:
        table = document.add_table(rows=1, cols=3)
        for index, label in enumerate(("Điều khoản", "Vấn đề", "Quyết định")):
            table.rows[0].cells[index].text = label
        for finding in outstanding:
            decision = by_key.get(str(finding.get("finding_key")), {})
            cells = table.add_row().cells
            cells[0].text = str(finding.get("clause", ""))
            cells[1].text = f"[{finding.get('severity', '')}] {finding.get('issue', '')}"
            cells[2].text = DECISION_LABELS.get(decision.get("decision"), "CHƯA QUYẾT ĐỊNH")
        style_docx_table(table, [1800, 5220, 2340])
    else:
        set_docx_font(
            document.add_paragraph().add_run("Mọi phát hiện đều đã được xử lý."),
            size=10, italic=True, color="667085",
        )

    clauses = review_result.get("clauses") or []
    if clauses:
        document.add_page_break()
        document.add_heading("PHỤ LỤC B — TOÀN VĂN ĐIỀU KHOẢN ĐÃ TRÍCH XUẤT", level=1)
        # The uploaded file is never stored, so the report carries the clause text it
        # refers to; otherwise the findings above would cite something unreadable.
        for clause in clauses:
            heading = document.add_paragraph()
            set_docx_font(
                heading.add_run(f"Điều {clause.get('number', '')}. {clause.get('title', '')}"),
                size=10, bold=True, color="1F4D78",
            )
            set_docx_font(document.add_paragraph().add_run(str(clause.get("text", ""))), size=9)

    notice = document.add_paragraph()
    notice.paragraph_format.space_before = Pt(12)
    notice.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_docx_font(
        notice.add_run(
            "LƯU Ý: Đây là báo cáo hỗ trợ đàm phán do AI tạo dựa trên quyết định của người rà "
            "soát. Legal phải xác nhận luật áp dụng và câu chữ cuối cùng trước khi ký."
        ),
        size=8, bold=True, color="9B1C1C",
    )

    output = io.BytesIO()
    document.save(output)
    return output.getvalue(), f"redline-{_safe_stem(document_name)}.docx"
