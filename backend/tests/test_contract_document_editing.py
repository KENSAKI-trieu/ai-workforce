"""Writing accepted revisions into the uploaded contract file, and sending it for approval.

The review's own clause list is what locates a clause in the file, so these tests build
it the way an upload does -- the file to Markdown, Markdown to clauses -- rather than by
hand, and the files themselves with python-docx and reportlab.
"""

import io
import json

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from pypdf import PdfReader

from app.core.config import settings
from app.domains.knowledge.document_markdown import to_markdown
from app.domains.legal.contract_document import (
    ContractEditError,
    apply_revisions,
    capture_form,
    decisions_fingerprint,
)
from app.domains.legal.contract_document.pdf_form import _font
from app.domains.legal.contract_review.clause_parser import split_contract_clauses


ARTICLES = [
    ("Điều 1. Phạm vi công việc", ["Bên B phát triển hệ thống quản lý kho cho Bên A."]),
    ("Điều 2. Thanh toán", ["Bên A thanh toán 100% trong vòng 90 ngày kể từ ngày ký."]),
    ("Điều 3. Phạt vi phạm", ["Bên vi phạm phải chịu phạt 30% giá trị hợp đồng."]),
    ("Điều 4. Chấm dứt", ["Bên A có quyền đơn phương chấm dứt bất kỳ lúc nào mà không cần bồi thường."]),
]


def _docx(*, table_in_article: int | None = None) -> bytes:
    document = Document()
    document.styles["Normal"].font.name = "Times New Roman"
    document.styles["Normal"].font.size = Pt(13)
    document.sections[0].header.paragraphs[0].text = "CÔNG TY ABC - HỢP ĐỒNG"
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("HỢP ĐỒNG PHÁT TRIỂN PHẦN MỀM").bold = True
    document.add_paragraph("Bên A: Công ty Khách hàng ABC")
    document.add_paragraph("Bên B: Công ty Phần mềm XYZ")
    for number, (heading, body) in enumerate(ARTICLES, 1):
        document.add_paragraph().add_run(heading).bold = True
        for line in body:
            document.add_paragraph(line).paragraph_format.first_line_indent = Pt(18)
        if number == table_in_article:
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "Đợt 1"
            table.cell(0, 1).text = "100%"
    signatures = document.add_table(rows=1, cols=2)
    signatures.cell(0, 0).text = "ĐẠI DIỆN BÊN A"
    signatures.cell(0, 1).text = "ĐẠI DIỆN BÊN B"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _pdf() -> bytes:
    from reportlab.pdfgen import canvas

    font = _font(True, False)
    bold = _font(True, True)
    buffer = io.BytesIO()
    page = canvas.Canvas(buffer, pagesize=(595, 842))
    y = 760
    page.setFont(bold, 14)
    page.drawCentredString(297, y, "HỢP ĐỒNG PHÁT TRIỂN PHẦN MỀM")
    y -= 40
    page.setFont(font, 12)
    for line in ("Bên A: Công ty Khách hàng ABC", "Bên B: Công ty Phần mềm XYZ"):
        page.drawString(72, y, line)
        y -= 22
    for heading, body in ARTICLES:
        y -= 10
        page.setFont(bold, 12)
        page.drawString(72, y, heading)
        y -= 18
        page.setFont(font, 12)
        for line in body:
            page.drawString(90, y, line)
            y -= 16
        y -= 30  # the blank space a revision may grow into
    page.setFont(bold, 12)
    page.drawString(72, y - 20, "ĐẠI DIỆN BÊN A")
    page.drawString(360, y - 20, "ĐẠI DIỆN BÊN B")
    page.showPage()
    page.save()
    return buffer.getvalue()


def _review(filename: str, data: bytes, revisions: dict[str, str], *, missing: str | None = None) -> dict:
    clauses = split_contract_clauses(to_markdown(filename, data))
    findings = []
    for number, text in revisions.items():
        clause = next(item for item in clauses if item["number"] == number)
        findings.append({
            "finding_key": f"k-{number}", "clause_id": clause["id"], "clause": number,
            "clause_title": clause["title"], "category": "OTHER", "finding_type": "LEGAL_ISSUE",
            "issue": f"Vấn đề điều {number}", "suggested_revision": text,
        })
    if missing:
        findings.append({
            "finding_key": "k-missing", "clause_id": None, "clause": "MISSING",
            "clause_title": "Điều khoản bị thiếu", "category": "CONFIDENTIALITY",
            "finding_type": "MISSING_CLAUSE", "issue": "Thiếu điều khoản: Bảo mật thông tin",
            "suggested_revision": missing,
        })
    return {
        "clauses": clauses,
        "findings": findings,
        "checklist": [{"category": "CONFIDENTIALITY", "label": "Bảo mật thông tin"}],
    }


def _accept_all(result: dict) -> list[dict]:
    return [{"finding_key": item["finding_key"], "decision": "ACCEPTED"} for item in result["findings"]]


PENALTY = "Bên vi phạm phải chịu phạt không vượt quá 8% giá trị phần nghĩa vụ bị vi phạm."
TERMINATION = (
    "Mỗi Bên chỉ được đơn phương chấm dứt khi Bên kia vi phạm nghiêm trọng và không khắc phục "
    "trong 15 ngày kể từ ngày nhận thông báo; Bên chấm dứt thanh toán phần việc đã thực hiện."
)
CONFIDENTIALITY = "Mỗi Bên bảo mật thông tin nhận được từ Bên kia trong và sau thời hạn hợp đồng."


# --------------------------------------------------------------------------- DOCX


def test_docx_revision_replaces_only_the_clause_body_and_keeps_its_form():
    data = _docx()
    result = _review("hd.docx", data, {"3": f"Điều 3. Phạt vi phạm\n{PENALTY}"}, missing=CONFIDENTIALITY)

    outcome = apply_revisions(
        data=data, filename="hd.docx", form=capture_form("hd.docx", data),
        review_result=result, decisions=_accept_all(result),
    )

    edited = Document(io.BytesIO(outcome["content"]))
    texts = [paragraph.text for paragraph in edited.paragraphs]
    # The heading is the document's own -- still bold, not repeated from the revision.
    assert texts.count("Điều 3. Phạt vi phạm") == 1
    heading = next(p for p in edited.paragraphs if p.text == "Điều 3. Phạt vi phạm")
    assert heading.runs[0].bold is True
    body = edited.paragraphs[texts.index("Điều 3. Phạt vi phạm") + 1]
    assert body.text == PENALTY
    assert body.paragraph_format.first_line_indent == Pt(18)
    assert "30% giá trị hợp đồng" not in "\n".join(texts)
    # A missing clause is added as the next article, before the signatures.
    assert "Điều 5. Bảo mật thông tin" in texts
    assert texts[texts.index("Điều 5. Bảo mật thông tin") + 1] == CONFIDENTIALITY
    assert edited.tables[-1].cell(0, 0).text == "ĐẠI DIỆN BÊN A"
    assert edited.sections[0].header.paragraphs[0].text == "CÔNG TY ABC - HỢP ĐỒNG"
    assert [item["action"] for item in outcome["report"]["applied"]] == ["REPLACED", "INSERTED"]
    assert all(check["ok"] for check in outcome["report"]["checks"])


def test_docx_clause_with_a_table_is_left_for_a_person():
    data = _docx(table_in_article=2)
    result = _review("hd.docx", data, {"2": "Bên A thanh toán trong 30 ngày.", "4": TERMINATION})

    outcome = apply_revisions(
        data=data, filename="hd.docx", form=capture_form("hd.docx", data),
        review_result=result, decisions=_accept_all(result),
    )

    assert [item["finding_key"] for item in outcome["report"]["skipped"]] == ["k-2"]
    assert "bảng" in outcome["report"]["skipped"][0]["reason"]
    text = to_markdown("hd.docx", outcome["content"])
    assert TERMINATION in text and "100% trong vòng 90 ngày" in text


def test_two_accepted_rewrites_of_one_clause_keep_the_reviewers_own_wording():
    data = _docx()
    result = _review("hd.docx", data, {"3": PENALTY})
    result["findings"].append({**result["findings"][0], "finding_key": "k-3b", "suggested_revision": "Không phạt."})
    decisions = [
        {"finding_key": "k-3", "decision": "ACCEPTED"},
        {"finding_key": "k-3b", "decision": "EDITED", "revised_text": "Mức phạt do hai Bên thoả thuận, tối đa 8%."},
    ]

    outcome = apply_revisions(data=data, filename="hd.docx", form=None, review_result=result, decisions=decisions)

    assert "Mức phạt do hai Bên thoả thuận, tối đa 8%." in to_markdown("hd.docx", outcome["content"])
    assert [item["finding_key"] for item in outcome["report"]["skipped"]] == ["k-3"]


def test_model_wording_on_one_line_becomes_the_clause_paragraphs():
    """As the model's revisions arrive live: heading, title and numbered points on one line."""
    data = _docx()
    result = _review("hd.docx", data, {
        "4": "Điều 4. Chấm dứt Mỗi Bên chỉ được chấm dứt khi Bên kia vi phạm nghiêm trọng.",
        "1": "Điều 1. Phạm vi công việc 1. Bên B phát triển hệ thống theo Phụ lục. 2. Yêu cầu phát sinh phải lập phụ lục.",
    }, missing="Điều [Số]. Giải quyết tranh chấp 1. Các Bên ưu tiên thương lượng. 2. Nếu không thành, đưa ra Tòa án.")
    result["findings"][-1]["issue"] = "Hợp đồng thiếu điều khoản giải quyết tranh chấp."

    outcome = apply_revisions(data=data, filename="hd.docx", form=None, review_result=result, decisions=_accept_all(result))

    texts = [paragraph.text for paragraph in Document(io.BytesIO(outcome["content"])).paragraphs]
    after = lambda heading, count: texts[texts.index(heading) + 1: texts.index(heading) + 1 + count]  # noqa: E731
    assert after("Điều 4. Chấm dứt", 1) == ["Mỗi Bên chỉ được chấm dứt khi Bên kia vi phạm nghiêm trọng."]
    assert after("Điều 1. Phạm vi công việc", 2) == [
        "1. Bên B phát triển hệ thống theo Phụ lục.", "2. Yêu cầu phát sinh phải lập phụ lục.",
    ]
    assert after("Điều 5. Giải quyết tranh chấp", 2) == [
        "1. Các Bên ưu tiên thương lượng.", "2. Nếu không thành, đưa ra Tòa án.",
    ]


def test_a_translated_review_is_not_written_into_the_original():
    data = _docx()
    result = {**_review("hd.docx", data, {"3": PENALTY}), "translated_for_review": True}
    with pytest.raises(ContractEditError):
        apply_revisions(data=data, filename="hd.docx", form=None, review_result=result, decisions=_accept_all(result))


def test_fingerprint_changes_with_the_accepted_wording_only():
    accepted = [{"finding_key": "a", "decision": "ACCEPTED"}, {"finding_key": "b", "decision": "REJECTED"}]
    assert decisions_fingerprint(accepted) == decisions_fingerprint(accepted[:1])
    assert decisions_fingerprint(accepted) != decisions_fingerprint(
        [{"finding_key": "a", "decision": "EDITED", "revised_text": "x"}]
    )


# --------------------------------------------------------------------------- PDF


def test_pdf_revision_removes_the_old_words_and_sets_the_new_in_place():
    data = _pdf()
    result = _review("hd.pdf", data, {"4": TERMINATION})

    outcome = apply_revisions(
        data=data, filename="hd.pdf", form=capture_form("hd.pdf", data),
        review_result=result, decisions=_accept_all(result),
    )

    reader = PdfReader(io.BytesIO(outcome["content"]))
    assert len(reader.pages) == 1
    text = reader.pages[0].extract_text()
    # Gone from the content stream, not painted over: it cannot be selected or found.
    assert "không cần bồi thường" not in text
    assert "vi phạm nghiêm trọng" in text
    assert "Điều 4. Chấm dứt" in text and "Điều 3. Phạt vi phạm" in text
    assert outcome["report"]["applied"][0]["action"] == "REPLACED"
    assert all(check["ok"] for check in outcome["report"]["checks"])


def test_pdf_revision_too_long_for_its_space_goes_to_an_amendment_annex():
    data = _pdf()
    long_text = " ".join([TERMINATION] * 4)
    result = _review("hd.pdf", data, {"4": long_text}, missing=CONFIDENTIALITY)

    outcome = apply_revisions(data=data, filename="hd.pdf", form=None, review_result=result, decisions=_accept_all(result))

    reader = PdfReader(io.BytesIO(outcome["content"]))
    assert len(reader.pages) == 2
    first, annex = reader.pages[0].extract_text(), reader.pages[1].extract_text()
    assert "không cần bồi thường" not in first
    assert "Phụ lục sửa đổi" in first
    assert "PHỤ LỤC SỬA ĐỔI, BỔ SUNG" in annex and "Bảo mật thông tin" in annex
    assert {item["action"] for item in outcome["report"]["applied"]} == {"ANNEXED"}
    assert outcome["report"]["warnings"]


# --------------------------------------------------------------------------- API


@pytest.fixture()
def legal_storage(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "LEGAL_DRAFT_STORAGE_PATH", str(tmp_path))
    return tmp_path


def _sse_events(body: str) -> list[tuple[str, dict]]:
    events = []
    for block in body.replace("\r\n", "\n").split("\n\n"):
        name, data = "message", []
        for line in block.split("\n"):
            if line.startswith("event:"):
                name = line[6:].strip()
            elif line.startswith("data:"):
                data.append(line[5:].strip())
        if data:
            events.append((name, json.loads("\n".join(data))))
    return events


def test_streamed_review_reports_each_stage_and_keeps_the_file(client, employee_token_headers, legal_storage):
    data = _docx()
    response = client.post(
        "/api/v1/legal/review-document/stream",
        files={"file": ("hop-dong-stream.docx", data, "application/octet-stream")},
        data={"represented_party": "PARTY_B"},
        headers=employee_token_headers,
    )
    assert response.status_code == 200
    events = _sse_events(response.text)
    stages = [payload["stage"] for name, payload in events if name == "progress"]
    for stage in ("UPLOAD", "FORM", "PARSE", "LANGUAGE", "SPLIT", "MAP", "SCORE", "SAVE"):
        assert stage in stages
    assert stages.index("FORM") < stages.index("PARSE")  # the form is taken before parsing
    name, review = events[-1]
    assert name == "complete"
    assert review["document"]["original_available"] is True
    assert review["document"]["original_format"] == "docx"
    assert review["document"]["editable"] is True
    stored = list(legal_storage.rglob("original/*"))
    assert stored and stored[0].read_bytes() != b""


def test_accept_apply_download_and_submit_for_approval(client, employee_token_headers, ceo_token_headers, legal_storage):
    data = _docx()
    review = client.post(
        "/api/v1/legal/review-document",
        files={"file": ("hop-dong-duyet.docx", data, "application/octet-stream")},
        # Not the side the streamed test used: the same text for the same side reopens that
        # review, whose file sits in the other test's storage.
        data={"represented_party": "PARTY_A"},
        headers=employee_token_headers,
    ).json()
    review_id = review["review_id"]
    target = next(item for item in review["findings"] if item["clause"] not in ("MISSING", "0"))
    client.put(
        f"/api/v1/legal/contract-reviews/{review_id}/decisions/{target['finding_key']}",
        json={"decision": "EDITED", "revised_text": "Nội dung điều khoản đã được thống nhất lại."},
        headers=employee_token_headers,
    ).raise_for_status()

    applied = client.post(
        f"/api/v1/legal/contract-reviews/{review_id}/revised-document", headers=employee_token_headers,
    )
    assert applied.status_code == 200, applied.text
    assert applied.json()["revised_ready"] is True and applied.json()["revised_stale"] is False

    download = client.get(f"/api/v1/legal/contract-reviews/{review_id}/revised-document", headers=employee_token_headers)
    assert download.status_code == 200
    assert "Nội dung điều khoản đã được thống nhất lại." in to_markdown("x.docx", download.content)
    original = client.get(f"/api/v1/legal/contract-reviews/{review_id}/original-document", headers=employee_token_headers)
    assert original.content == data

    # A new decision makes the file stale until it is written again.
    other = next(item for item in review["findings"] if item["finding_key"] != target["finding_key"])
    client.put(
        f"/api/v1/legal/contract-reviews/{review_id}/decisions/{other['finding_key']}",
        json={"decision": "REJECTED"}, headers=employee_token_headers,
    ).raise_for_status()
    reopened = client.get(f"/api/v1/legal/contract-reviews/{review_id}", headers=employee_token_headers).json()
    assert reopened["document"]["revised_stale"] is False  # a rejection changes no wording

    submitted = client.post(
        f"/api/v1/legal/contract-reviews/{review_id}/submit-approval",
        json={"note": "Nhờ anh chị duyệt"}, headers=employee_token_headers,
    )
    assert submitted.status_code == 200, submitted.text
    body = submitted.json()
    assert body["status"] in {"CREATED", "UPDATED"}
    assert body["approval"]["status"] == "WAITING"
    assert body["approval"]["eligible_approver_count"] >= 1
    assert body["approval"]["warning"] is None

    # The requester may not sign it, so it is not in their queue -- it is in what they sent.
    own_queue = client.get("/api/v1/approvals/pending", headers=employee_token_headers).json()
    assert all(item["payload"].get("contract_review_id") != review_id for item in own_queue)
    sent = client.get("/api/v1/approvals/submitted", headers=employee_token_headers).json()
    mine = next(item for item in sent if item["payload"].get("contract_review_id") == review_id)
    assert mine["status"] == "WAITING" and mine["eligible_approver_count"] >= 1 and mine["warning"] is None

    pending = client.get("/api/v1/approvals/pending", headers=ceo_token_headers).json()
    card = next(item for item in pending if item["payload"].get("contract_review_id") == review_id)
    assert card["payload"]["submitted_manually"] is True
    assert card["payload"]["revised_document"]["url"].endswith(f"{review_id}/revised-document")
    assert card["payload"]["note"] == "Nhờ anh chị duyệt"
    assert "original_text" not in json.dumps(card["payload"])  # no contract wording in the payload

    again = client.post(
        f"/api/v1/legal/contract-reviews/{review_id}/submit-approval", json={}, headers=employee_token_headers,
    ).json()
    assert again["status"] == "UPDATED"
    assert again["approval"]["approval_id"] == body["approval"]["approval_id"]


def test_a_review_without_its_file_cannot_be_edited(client, employee_token_headers, legal_storage):
    review = client.post(
        "/api/v1/legal/review-document",
        files={"file": ("hop-dong.txt", "HỢP ĐỒNG\nĐiều 1. Phạt\nPhạt 30% giá trị hợp đồng.\n".encode(), "text/plain")},
        data={"represented_party": "PARTY_A"},
        headers=employee_token_headers,
    ).json()
    assert review["document"]["original_available"] is False
    response = client.post(
        f"/api/v1/legal/contract-reviews/{review['review_id']}/revised-document", headers=employee_token_headers,
    )
    assert response.status_code == 409
