"""Deleting a saved contract review: its creator's call, never under a waiting approval."""

import io
import uuid

import pytest
from docx import Document

from app.core.config import settings
from app.models.models import ContractReview, ContractReviewDecision


def _docx(tag: str) -> bytes:
    document = Document()
    document.add_paragraph(f"HỢP ĐỒNG DỊCH VỤ {tag}")
    document.add_paragraph("Điều 1. Phạm vi")
    document.add_paragraph("Bên B cung cấp dịch vụ vận hành hệ thống cho Bên A.")
    document.add_paragraph("Điều 2. Phạt vi phạm")
    document.add_paragraph("Bên vi phạm phải chịu phạt 30% giá trị hợp đồng.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest.fixture()
def legal_storage(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "LEGAL_DRAFT_STORAGE_PATH", str(tmp_path))
    return tmp_path


def _review(client, headers, tag):
    response = client.post(
        "/api/v1/legal/review-document",
        files={"file": (f"hop-dong-xoa-{tag}.docx", _docx(tag), "application/octet-stream")},
        data={"represented_party": "PARTY_A"},
        headers=headers,
    )
    assert response.status_code == 200, response.text
    return response.json()


def test_the_creator_deletes_a_review_with_its_decisions_and_files(
    client, employee_token_headers, transactional_db_session, legal_storage
):
    review = _review(client, employee_token_headers, uuid.uuid4().hex[:8])
    review_id = review["review_id"]
    assert review["can_delete"] is True
    finding = next(item for item in review["findings"] if item["clause"] not in ("MISSING", "0"))
    client.put(
        f"/api/v1/legal/contract-reviews/{review_id}/decisions/{finding['finding_key']}",
        json={"decision": "EDITED", "revised_text": "Mức phạt không vượt quá 8%."},
        headers=employee_token_headers,
    ).raise_for_status()
    client.post(
        f"/api/v1/legal/contract-reviews/{review_id}/revised-document", headers=employee_token_headers,
    ).raise_for_status()
    assert any(path.is_file() for path in legal_storage.rglob("*"))

    response = client.delete(f"/api/v1/legal/contract-reviews/{review_id}", headers=employee_token_headers)

    assert response.status_code == 200, response.text
    assert client.get(f"/api/v1/legal/contract-reviews/{review_id}", headers=employee_token_headers).status_code == 404
    listed = client.get("/api/v1/legal/contract-reviews", headers=employee_token_headers).json()
    assert all(item["review_id"] != review_id for item in listed)
    transactional_db_session.expire_all()
    assert transactional_db_session.query(ContractReviewDecision).filter(
        ContractReviewDecision.review_id == uuid.UUID(review_id)
    ).count() == 0
    # The original and the revised file go with it.
    assert not [path for path in legal_storage.rglob("*") if path.is_file()]


def test_a_review_waiting_on_an_approval_is_not_deleted(client, employee_token_headers, legal_storage):
    review = _review(client, employee_token_headers, uuid.uuid4().hex[:8])
    client.post(
        f"/api/v1/legal/contract-reviews/{review['review_id']}/submit-approval", json={}, headers=employee_token_headers,
    ).raise_for_status()
    reopened = client.get(f"/api/v1/legal/contract-reviews/{review['review_id']}", headers=employee_token_headers).json()
    assert reopened["can_delete"] is False

    response = client.delete(f"/api/v1/legal/contract-reviews/{review['review_id']}", headers=employee_token_headers)

    assert response.status_code == 409
    assert client.get(f"/api/v1/legal/contract-reviews/{review['review_id']}", headers=employee_token_headers).status_code == 200


def test_only_the_creator_may_delete(client, employee_token_headers, ceo_token_headers, legal_storage):
    review = _review(client, employee_token_headers, uuid.uuid4().hex[:8])
    url = f"/api/v1/legal/contract-reviews/{review['review_id']}"
    seen_by_ceo = client.get(url, headers=ceo_token_headers)
    if seen_by_ceo.status_code == 200:
        # Allowed to read every review is not allowed to remove one.
        assert seen_by_ceo.json()["can_delete"] is False
        assert client.delete(url, headers=ceo_token_headers).status_code == 403
    else:
        assert client.delete(url, headers=ceo_token_headers).status_code == 404
    assert client.get(url, headers=employee_token_headers).status_code == 200


def test_a_later_round_outlives_the_review_it_followed(
    client, employee_token_headers, transactional_db_session, legal_storage
):
    first = _review(client, employee_token_headers, uuid.uuid4().hex[:8])
    second = _review(client, employee_token_headers, uuid.uuid4().hex[:8])
    row = transactional_db_session.get(ContractReview, uuid.UUID(second["review_id"]))
    row.parent_review_id = uuid.UUID(first["review_id"])
    transactional_db_session.commit()

    assert client.delete(
        f"/api/v1/legal/contract-reviews/{first['review_id']}", headers=employee_token_headers,
    ).status_code == 200

    transactional_db_session.expire_all()
    later = transactional_db_session.get(ContractReview, uuid.UUID(second["review_id"]))
    assert later is not None and later.parent_review_id is None
