"""The next round of a contract review: the revised file checked against the round before.

Live 2026-09-30: a reviewer revised a contract as the review proposed, uploaded the revised
file, and the fresh review faulted -- at medium and low -- the very wording the round before
had written; round after round, the findings contradicted each other. A revised file now
carries a mark naming its review, and uploaded again it is read against that review: an
untouched clause keeps its finding and decision, an accepted risk is not raised again, and
for a changed clause the model says whether the change fixed what was found.
"""

import io
import uuid
import json

import pytest
from docx import Document

from app.core.config import settings
from app.domains.legal import contract_review_store
from app.domains.legal.contract_document import read_review_marker, stamp_review_marker
from app.domains.legal.contract_review.llm_assessment import (
    MAP_SYSTEM_PROMPT,
    REVIEW_SYSTEM_PROMPT,
)
from app.domains.legal.contract_review.rereview import VERIFY_SYSTEM_PROMPT, match_clauses
from app.models.models import ContractReview, WorkflowApproval

ARTICLES = [
    ("Điều 1. Phạm vi công việc", "Bên B phát triển hệ thống quản lý kho cho Bên A."),
    ("Điều 2. Thanh toán", "Bên A thanh toán 100% trong vòng 90 ngày kể từ ngày nghiệm thu."),
    ("Điều 3. Chấm dứt", "Bên A có quyền đơn phương chấm dứt bất kỳ lúc nào mà không cần bồi thường."),
    ("Điều 4. Bảo hành", "Bên B bảo hành sản phẩm trong 3 tháng kể từ ngày nghiệm thu."),
]
CATEGORY = {"1": "SCOPE", "2": "PAYMENT", "3": "TERMINATION", "4": "WARRANTY"}
REVIEWED = {
    "2": ("MEDIUM", "Thời hạn thanh toán 90 ngày quá dài"),
    "3": ("HIGH", "Bên A được đơn phương chấm dứt không bồi thường"),
    "4": ("LOW", "Thời hạn bảo hành ngắn"),
}
TERMINATION_FIX = "Mỗi Bên được chấm dứt khi Bên kia vi phạm nghiêm trọng và không khắc phục trong 15 ngày."


def _docx() -> bytes:
    document = Document()
    document.add_paragraph("HỢP ĐỒNG PHÁT TRIỂN PHẦN MỀM")
    document.add_paragraph("Bên A: Công ty Khách hàng ABC")
    document.add_paragraph("Bên B: Công ty Phần mềm XYZ. Người đại diện: Ông Nguyễn Văn An, điện thoại 0912 345 678")
    for heading, body in ARTICLES:
        document.add_paragraph(heading)
        document.add_paragraph(body)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class RoundModel:
    """Reads a first round (map, review, revise) and judges the next one (verify)."""

    enabled = True

    def __init__(self):
        self.calls: list[tuple[str, dict]] = []

    def generate_text(self, messages, **_kwargs):
        system, payload = messages[0]["content"], json.loads(messages[1]["content"])
        stage = {MAP_SYSTEM_PROMPT: "map", REVIEW_SYSTEM_PROMPT: "review", VERIFY_SYSTEM_PROMPT: "verify"}.get(system, "revise")
        self.calls.append((stage, payload))
        if stage == "map":
            reply = {
                "contract_type": "SOFTWARE_DEVELOPMENT_CONTRACT", "contract_type_confidence": 0.95,
                "parties": {"PARTY_A": {"name": "Công ty Khách hàng ABC", "role": "Bên đặt hàng"},
                            "PARTY_B": {"name": "Công ty Phần mềm XYZ", "role": "Bên phát triển"}},
                "clause_categories": [
                    {"clause_id": clause["id"], "category": CATEGORY.get(clause["number"], "PARTIES")}
                    for clause in payload["clauses"]
                ],
                "checklist": [], "findings": [],
            }
        elif stage == "review":
            reply = {"findings": [
                {"clause_id": clause["id"], "category": clause["category"], "finding_type": "COMMERCIAL_RISK",
                 "severity": REVIEWED[clause["number"]][0], "impact": "ADVERSE", "confidence": 0.8,
                 "issue": REVIEWED[clause["number"]][1], "evidence": "", "reason": "Bất lợi cho Bên B.",
                 "legal_basis": None, "recommendation": "Đàm phán lại."}
                for clause in payload["clauses_to_review"] if clause["number"] in REVIEWED
            ]}
        elif stage == "verify":
            changed = {clause["id"] for clause in payload["clauses"] if clause["status"] == "CHANGED"}
            reply = {
                "verdicts": [
                    {"id": problem["id"], "status": "RESOLVED" if problem["fix_applied"] else "UNRESOLVED",
                     "severity": problem["severity"], "clause_id": problem["clause_id"],
                     "note": "Đã sửa theo đề xuất." if problem["fix_applied"] else "Chưa sửa."}
                    for problem in payload["problems"]
                ],
                "new_findings": [
                    {"clause_id": clause_id, "category": "TERMINATION", "finding_type": "AMBIGUOUS_CLAUSE",
                     "severity": "LOW", "impact": "SHARED", "confidence": 0.6,
                     "issue": "Chưa định nghĩa vi phạm nghiêm trọng", "evidence": "vi phạm nghiêm trọng",
                     "reason": "Dễ tranh cãi.", "legal_basis": None, "recommendation": "Định nghĩa."}
                    for clause_id in sorted(changed)
                ],
            }
        else:
            reply = {"revisions": [
                {"id": problem["id"], "suggested_revision": f"Đề xuất cho: {problem['issue']}"}
                for problem in payload["problems"]
            ]}
        return {"provider": "gemini", "content": json.dumps(reply, ensure_ascii=False), "usage": {"total_tokens": 5}}


@pytest.fixture()
def model(monkeypatch):
    fake = RoundModel()
    monkeypatch.setattr("app.domains.legal.contract_review.llm_assessment.get_ai_service_client", lambda: fake)
    return fake


@pytest.fixture()
def legal_storage(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "LEGAL_DRAFT_STORAGE_PATH", str(tmp_path))
    return tmp_path


def _upload(client, headers, filename, data, side="PARTY_B"):
    response = client.post(
        "/api/v1/legal/review-document",
        files={"file": (filename, data, "application/octet-stream")},
        data={"represented_party": side},
        headers=headers,
    )
    assert response.status_code == 200, response.text
    return response.json()


def _finding(review, issue):
    return next(item for item in review["findings"] if item["issue"] == issue)


def _decide(client, headers, review, finding, decision, text=None):
    client.put(
        f"/api/v1/legal/contract-reviews/{review['review_id']}/decisions/{finding['finding_key']}",
        json={"decision": decision, "revised_text": text}, headers=headers,
    ).raise_for_status()


def test_a_revised_file_is_checked_against_the_round_it_came_from(
    client, employee_token_headers, transactional_db_session, legal_storage, model
):
    first = _upload(client, employee_token_headers, "hop-dong-vong.docx", _docx())
    assert first["review_engine"] == "LLM_ASSISTED" and first["review_round"] == 1
    # High risk, and still nothing sent: the reviewer sends it when ready.
    assert first["requires_legal_approval"] is True and first["approval"] is None
    termination = _finding(first, REVIEWED["3"][1])
    warranty = _finding(first, REVIEWED["4"][1])
    payment = _finding(first, REVIEWED["2"][1])
    _decide(client, employee_token_headers, first, termination, "EDITED", TERMINATION_FIX)
    _decide(client, employee_token_headers, first, warranty, "REJECTED")  # risk accepted
    client.post(
        f"/api/v1/legal/contract-reviews/{first['review_id']}/revised-document", headers=employee_token_headers,
    ).raise_for_status()
    revised = client.get(
        f"/api/v1/legal/contract-reviews/{first['review_id']}/revised-document", headers=employee_token_headers,
    ).content
    assert read_review_marker("docx", revised).startswith(first["review_id"])

    model.calls.clear()
    second = _upload(client, employee_token_headers, "hop-dong-vong-da-sua.docx", revised)

    assert second["review_round"] == 2 and second["parent_review_id"] == first["review_id"]
    assert second["round_mode"] == "VERIFIED"
    # Read against the round before: no map, no clause-by-clause review.
    assert [stage for stage, _ in model.calls] == ["verify", "revise"]
    verify = model.calls[0][1]
    problems = {problem["issue"]: problem for problem in verify["problems"]}
    assert problems[REVIEWED["3"][1]]["fix_applied"] is True
    assert TERMINATION_FIX in problems[REVIEWED["3"][1]]["proposed_fix"]
    # Neither the untouched clause nor the accepted risk is put to the model again.
    assert REVIEWED["2"][1] not in problems and REVIEWED["4"][1] not in problems
    reported = {item["issue"]: item for item in verify["already_reported"]}
    assert reported[REVIEWED["4"][1]]["accepted_by_reviewer"] is True
    assert reported[REVIEWED["2"][1]]["accepted_by_reviewer"] is False
    sent = json.dumps(model.calls, ensure_ascii=False)
    assert "0912 345 678" not in sent and "Nguyễn Văn An" not in sent

    # The fix settled the problem it was written for.
    assert REVIEWED["3"][1] in {item["issue"] for item in second["resolved_findings"]}
    assert all(item["issue"] != REVIEWED["3"][1] for item in second["findings"])
    # The untouched clause keeps its finding word for word, wording proposed included.
    kept = _finding(second, REVIEWED["2"][1])
    assert kept["round_status"] == "CARRIED"
    assert kept["suggested_revision"] == payment["suggested_revision"]
    # The accepted risk comes back decided, not as a question.
    accepted = _finding(second, REVIEWED["4"][1])
    assert accepted["round_status"] == "ACCEPTED_RISK"
    decisions = {item["finding_key"]: item for item in second["decisions"]}
    assert decisions[accepted["finding_key"]]["decision"] == "REJECTED"
    assert "Giữ từ bản rà soát" in decisions[accepted["finding_key"]]["comment"]
    # What the change itself brought in is new, with wording of its own.
    new = _finding(second, "Chưa định nghĩa vi phạm nghiêm trọng")
    assert new["round_status"] == "NEW" and new["suggested_revision"].startswith("Đề xuất cho")
    assert second["round_summary"]["RESOLVED"] >= 1 and second["round_summary"]["NEW"] >= 1

    stored = transactional_db_session.get(ContractReview, uuid.UUID(second["review_id"]))
    assert str(stored.parent_review_id) == first["review_id"]
    assert all(
        (approval.payload or {}).get("contract_review_id") not in {first["review_id"], second["review_id"]}
        for approval in transactional_db_session.query(WorkflowApproval).all()
    )


def test_a_mark_from_someone_elses_review_or_a_forged_one_is_ignored(
    client, employee_token_headers, legal_storage, model
):
    data = _docx()
    forged = stamp_review_marker("docx", data, "00000000-0000-0000-0000-000000000000.deadbeef")
    review = _upload(client, employee_token_headers, "hop-dong-gia.docx", forged, side="PARTY_A")
    assert review["review_round"] == 1 and review["parent_review_id"] is None


def test_a_file_that_is_no_longer_the_same_contract_is_reviewed_in_full(
    client, employee_token_headers, transactional_db_session, legal_storage, model
):
    first = _upload(client, employee_token_headers, "hop-dong-goc.docx", _docx(), side="NEUTRAL")
    row = transactional_db_session.get(ContractReview, uuid.UUID(first["review_id"]))
    document = Document()
    document.add_paragraph("HỢP ĐỒNG MUA BÁN HÀNG HÓA")
    for number in range(1, 6):
        document.add_paragraph(f"Điều {number}. Nội dung {number}")
        document.add_paragraph(f"Bên bán giao lô hàng số {number} tại kho của bên mua trong tháng {number}.")
    buffer = io.BytesIO()
    document.save(buffer)
    # A genuine mark on a different contract: not the next round of anything.
    marked = stamp_review_marker("docx", buffer.getvalue(), contract_review_store.review_marker(row))

    other = _upload(client, employee_token_headers, "hop-dong-khac.docx", marked, side="NEUTRAL")

    assert other["review_round"] == 1 and other["parent_review_id"] is None
    assert "verify" not in [stage for stage, _ in model.calls]


def test_clauses_are_paired_across_a_renumbering():
    old = [
        {"id": "clause-1", "number": "1", "title": "Phạm vi", "text": "Điều 1. Phạm vi\nBên B phát triển hệ thống quản lý kho."},
        {"id": "clause-2", "number": "2", "title": "Thanh toán", "text": "Điều 2. Thanh toán\nBên A thanh toán trong 90 ngày kể từ ngày nghiệm thu."},
        {"id": "clause-3", "number": "3", "title": "Phạt", "text": "Điều 3. Phạt\nBên vi phạm chịu phạt 30% giá trị hợp đồng."},
    ]
    new = [
        {"id": "clause-1", "number": "1", "title": "Phạm vi", "text": "Điều 1. Phạm vi\nBên B phát triển hệ thống quản lý kho."},
        {"id": "clause-2", "number": "2", "title": "Bảo mật", "text": "Điều 2. Bảo mật\nCác Bên giữ bí mật thông tin của nhau."},
        {"id": "clause-3", "number": "3", "title": "Thanh toán", "text": "Điều 3. Thanh toán\nBên A thanh toán trong 90 ngày kể từ ngày nghiệm thu."},
        {"id": "clause-4", "number": "4", "title": "Phạt", "text": "Điều 4. Phạt\nBên vi phạm chịu phạt 8% giá trị phần nghĩa vụ bị vi phạm."},
    ]
    match = match_clauses(old, new)
    # Renumbered but untouched: the same clause, unchanged.
    assert match.new_for_old["clause-2"]["id"] == "clause-3" and "clause-2" not in match.changed_old
    # Same title, new terms: the same clause, changed.
    assert match.new_for_old["clause-3"]["id"] == "clause-4" and "clause-3" in match.changed_old
    assert [clause["id"] for clause in match.added] == ["clause-2"]
    assert match.removed == []


def test_the_mark_survives_in_word_and_pdf():
    data = _docx()
    assert read_review_marker("docx", data) is None
    stamped = stamp_review_marker("docx", data, "abc.123")
    assert read_review_marker("docx", stamped) == "abc.123"
    # Stamped again, the mark is replaced, not doubled; the document still opens.
    again = stamp_review_marker("docx", stamped, "def.456")
    assert read_review_marker("docx", again) == "def.456"
    assert Document(io.BytesIO(again)).paragraphs[0].text == "HỢP ĐỒNG PHÁT TRIỂN PHẦN MỀM"

    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    page = canvas.Canvas(buffer)
    page.drawString(72, 720, "Contract")
    page.save()
    pdf = stamp_review_marker("pdf", buffer.getvalue(), "abc.123")
    assert read_review_marker("pdf", pdf) == "abc.123"
    assert read_review_marker("pdf", b"not a pdf") is None
