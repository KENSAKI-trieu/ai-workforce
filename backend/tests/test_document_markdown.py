"""PDF and DOCX reach the contract pipeline as Markdown, structure intact."""

import io
import os

import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

from app.domains.knowledge.document_markdown import (
    docx_to_markdown,
    extract_knowledge_text,
    pdf_to_markdown,
    to_markdown,
)
from app.domains.knowledge.rag_chunking import chunk_text
from app.domains.legal.contract_review.clause_parser import split_contract_clauses

# A font with Vietnamese marks for the generated PDFs: Windows locally, DejaVu on the CI's
# Ubuntu runner. reportlab's own Vera has no "ạ".
FONT = next(
    (path for path in (
        "C:/Windows/Fonts/arial.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/dejavu/DejaVuSans.ttf",
        "/Library/Fonts/Arial Unicode.ttf",
    ) if os.path.exists(path)),
    None,
)
needs_font = pytest.mark.skipif(FONT is None, reason="no font with Vietnamese glyphs to build a PDF")


def _auto_numbered_docx() -> bytes:
    """Articles numbered by Word ("Điều %1."), points under them, a table between two."""
    doc = Document()
    numbering = doc.part.numbering_part.element
    numbering.append(parse_xml(
        f'<w:abstractNum {nsdecls("w")} w:abstractNumId="7">'
        '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="Điều %1."/></w:lvl>'
        '<w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="lowerLetter"/><w:lvlText w:val="%2)"/></w:lvl>'
        '</w:abstractNum>'))
    numbering.append(parse_xml(f'<w:num {nsdecls("w")} w:numId="9"><w:abstractNumId w:val="7"/></w:num>'))

    def numbered(text, level=0):
        paragraph = doc.add_paragraph(text)
        paragraph._p.get_or_add_pPr().append(parse_xml(
            f'<w:numPr {nsdecls("w")}><w:ilvl w:val="{level}"/><w:numId w:val="9"/></w:numPr>'))

    doc.add_paragraph("HỢP ĐỒNG MUA BÁN HÀNG HÓA")
    doc.add_paragraph("Bên A: Công ty TNHH Hưng Thịnh\nĐịa chỉ: KCN Quang Minh")
    numbered("Hàng hóa")
    table = doc.add_table(rows=3, cols=3)
    for r, row in enumerate([("STT", "Tên hàng", "Thành tiền"), ("1", "Máy phát | 500kVA", "925.000.000"), ("2", "Tủ ATS", "120.000.000")]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    merged = doc.add_table(rows=1, cols=3)
    merged.cell(0, 0).merge(merged.cell(0, 1)).text = "Tổng cộng"
    merged.cell(0, 2).text = "1.045.000.000"
    numbered("Thanh toán")
    numbered("Tạm ứng 70% trong 03 ngày.", 1)
    numbered("Phần còn lại trong 05 ngày.", 1)
    numbered("Phạt vi phạm")
    numbered("Phạt 15% giá trị hợp đồng.", 1)
    doc.add_heading("Phụ lục", level=2)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def test_a_docx_keeps_words_article_numbers_and_its_tables_in_place():
    markdown = docx_to_markdown(_auto_numbered_docx())

    assert markdown.splitlines()[0] == "# HỢP ĐỒNG MUA BÁN HÀNG HÓA"
    # A soft line break keeps the address on a line of its own.
    assert "Bên A: Công ty TNHH Hưng Thịnh\nĐịa chỉ: KCN Quang Minh" in markdown
    order = [markdown.index(part) for part in (
        "### Điều 1. Hàng hóa", "| STT | Tên hàng | Thành tiền |", "| Tổng cộng | 1.045.000.000 |",
        "### Điều 2. Thanh toán", "a) Tạm ứng 70%", "b) Phần còn lại", "### Điều 3. Phạt vi phạm",
        "a) Phạt 15%", "## Phụ lục",
    )]
    assert order == sorted(order)
    assert "| 1 | Máy phát \\| 500kVA | 925.000.000 |" in markdown  # a pipe in a cell is escaped
    assert "|\n| --- | --- | --- |\n" in markdown


def test_word_numbering_restarts_the_points_under_each_article():
    markdown = docx_to_markdown(_auto_numbered_docx())

    under_article_3 = markdown[markdown.index("### Điều 3."):]
    assert "a) Phạt 15%" in under_article_3 and "c)" not in under_article_3


def _pdf(pages: list[list[tuple[float, float, str]]]) -> bytes:
    pdfmetrics.registerFont(TTFont("Arial", FONT))
    out = io.BytesIO()
    pdf = canvas.Canvas(out, pagesize=A4)
    for number, items in enumerate(pages, 1):
        pdf.setFont("Arial", 9)
        pdf.drawString(50, 810, "CÔNG TY TNHH HƯNG THỊNH - HĐ số 15/2026/HĐMB")
        pdf.drawString(280, 30, f"Trang {number}/{len(pages)}")
        pdf.setFont("Arial", 11)
        for x, y, text in items:
            pdf.drawString(x, y, text)
        pdf.showPage()
    pdf.save()
    return out.getvalue()


LONG = "Bên Mua tạm ứng 70% giá trị hợp đồng trong vòng 03 ngày kể từ ngày ký hợp đồng này, phần"


@pytest.fixture(scope="module")
def contract_pdf() -> bytes:
    return _pdf([
        [
            (50, 770, "HỢP ĐỒNG MUA BÁN HÀNG HÓA"),
            (50, 740, "Điều 1. Hàng hóa và giá cả"),
            (50, 716, "Tên hàng"), (260, 716, "Thành tiền"),
            (50, 692, "Máy phát điện công nghiệp"), (260, 692, "925.000.000"),
            (50, 678, "500kVA, xuất xứ Nhật Bản"),
            (50, 654, "Tủ ATS"), (260, 654, "120.000.000"),
            (50, 624, "Điều 2. Thanh toán"),
            (50, 600, LONG),
        ],
        [
            (50, 770, "còn lại thanh toán trong vòng 05 ngày sau khi nhận hàng và hóa đơn hợp lệ."),
            (50, 740, "Điều 3. Nghĩa vụ của Bên Bán"),
            (60, 716, "•    Giao hàng đúng hạn."),
            (60, 698, "•    Bảo hành 06 tháng."),
        ],
        [
            (50, 770, "Điều 4. Giải quyết tranh chấp"),
            (50, 746, "Mọi tranh chấp được giải quyết tại SIAC."),
        ],
    ])


@needs_font
def test_a_pdf_drops_page_headers_and_numbers(contract_pdf):
    markdown = pdf_to_markdown(contract_pdf)

    assert "CÔNG TY TNHH HƯNG THỊNH - HĐ số" not in markdown
    assert "Trang " not in markdown


@needs_font
def test_a_pdf_rejoins_a_sentence_the_page_break_cut(contract_pdf):
    markdown = pdf_to_markdown(contract_pdf)

    assert f"{LONG} còn lại thanh toán trong vòng 05 ngày" in markdown


@needs_font
def test_a_pdf_table_comes_back_as_rows_with_wrapped_cells_joined(contract_pdf):
    markdown = pdf_to_markdown(contract_pdf)

    assert "| Tên hàng | Thành tiền |\n| --- | --- |" in markdown
    assert "| Máy phát điện công nghiệp 500kVA, xuất xứ Nhật Bản | 925.000.000 |" in markdown
    assert "| Tủ ATS | 120.000.000 |" in markdown


@needs_font
def test_a_pdf_marks_articles_and_lists(contract_pdf):
    markdown = pdf_to_markdown(contract_pdf)

    assert markdown.startswith("# HỢP ĐỒNG MUA BÁN HÀNG HÓA")
    for number in range(1, 5):
        assert f"\n### Điều {number}." in markdown
    assert "- Giao hàng đúng hạn.\n- Bảo hành 06 tháng." in markdown


def test_a_scanned_pdf_has_no_text_to_give():
    from PIL import Image

    scan = io.BytesIO()
    Image.new("RGB", (400, 200), "white").save(scan, format="PDF")
    assert to_markdown("scan.pdf", scan.getvalue()) == ""


def test_other_files_are_left_to_the_plain_parser():
    assert to_markdown("notes.txt", b"hello") is None


@needs_font
def test_the_clause_splitter_reads_the_markdown(contract_pdf):
    clauses = split_contract_clauses(pdf_to_markdown(contract_pdf))

    articles = [(clause["number"], clause["title"]) for clause in clauses if clause["number"] != "0"]
    assert articles == [("1", "Hàng hóa và giá cả"), ("2", "Thanh toán"), ("3", "Nghĩa vụ của Bên Bán"), ("4", "Giải quyết tranh chấp")]
    assert "| Tủ ATS | 120.000.000 |" in clauses[1]["text"]


def test_points_stay_inside_their_article():
    text = "Điều 1. Hàng hóa\n1. Máy phát điện.\n2. Tủ ATS.\nĐiều 2. Thanh toán\n1. Tạm ứng 70%.\nPHỤ LỤC GIÁ"

    clauses = split_contract_clauses(text)

    assert [clause["number"] for clause in clauses] == ["1", "2"]
    assert "2. Tủ ATS." in clauses[0]["text"] and "PHỤ LỤC GIÁ" in clauses[1]["text"]


def test_a_document_without_articles_still_splits_at_its_numbers():
    clauses = split_contract_clauses("1. Phạm vi\nCung cấp phần mềm.\n2. Thanh toán\nTrong 30 ngày.")

    assert [(clause["number"], clause["title"]) for clause in clauses] == [("1", "Phạm vi"), ("2", "Thanh toán")]


# ------------------------------------------------------------------ knowledge index


@needs_font
def test_page_markers_open_each_page_for_the_knowledge_index(contract_pdf):
    markdown = pdf_to_markdown(contract_pdf, page_markers=True)

    assert markdown.startswith("[[PAGE:1]]\n# ")
    assert "[[PAGE:2]]\n### Điều 3." in markdown
    assert "[[PAGE:3]]\n### Điều 4." in markdown
    # The sentence the page break cut stays whole, counted on the page it starts on.
    assert f"{LONG} còn lại" in markdown
    assert "[[PAGE:" not in pdf_to_markdown(contract_pdf)


@needs_font
def test_the_knowledge_index_cites_the_page_of_each_article(contract_pdf):
    chunks = chunk_text(
        extract_knowledge_text("contract.pdf", contract_pdf),
        chunk_size=700, chunk_overlap=80, target_size=450, min_size=0,
    )

    pages = {chunk["section_title"]: chunk["page"] for chunk in chunks}
    assert pages["Điều 1. Hàng hóa và giá cả"] == 1
    assert pages["Điều 3. Nghĩa vụ của Bên Bán"] == 2
    assert pages["Điều 4. Giải quyết tranh chấp"] == 3


def test_the_knowledge_index_chunks_a_docx_at_the_articles_word_numbered():
    chunks = chunk_text(
        extract_knowledge_text("contract.docx", _auto_numbered_docx()),
        chunk_size=700, chunk_overlap=80, target_size=450, min_size=0,
    )

    titles = [chunk["section_title"] for chunk in chunks]
    assert {"Điều 1. Hàng hóa", "Điều 2. Thanh toán", "Điều 3. Phạt vi phạm"} <= set(titles)
    article_1 = next(chunk for chunk in chunks if chunk["section_title"] == "Điều 1. Hàng hóa")
    assert "| Tủ ATS | 120.000.000 |" in article_1["content"]  # the table stays in its article


def test_a_scan_gives_the_knowledge_index_no_text():
    from PIL import Image

    scan = io.BytesIO()
    Image.new("RGB", (400, 200), "white").save(scan, format="PDF")
    assert extract_knowledge_text("scan.pdf", scan.getvalue()) == ""
    assert extract_knowledge_text("notes.txt", "Ghi chú".encode()) == "Ghi chú"
