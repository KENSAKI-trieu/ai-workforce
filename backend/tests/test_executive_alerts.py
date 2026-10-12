"""The CEO assistant's C4: proactive alerts, the morning digest, and its own page.

What is pinned: a rule finds only what the recipient's position may see; thresholds are
checked, not clipped; the scheduler sends the digest once per chosen day and never
repeats an alert the same day; and the page's endpoints keep the two permission boxes.
"""

from __future__ import annotations

import uuid
from datetime import date, datetime, timedelta, timezone
from zoneinfo import ZoneInfo

import pytest
from fastapi import HTTPException

from app.domains.executive import alerts as alerts_module
from app.domains.executive import scheduler
from app.domains.executive.alerts import RULES, evaluate_alerts, rule_settings, save_rule_settings
from app.domains.executive.digest import digest_lines
from app.models.models import ExecutiveSchedule, Notification, Task, User
from tests.finance_helpers import login, person

VN = ZoneInfo("Asia/Ho_Chi_Minh")


@pytest.fixture()
def db(transactional_db_session):
    return transactional_db_session


@pytest.fixture()
def admin(db):
    return db.query(User).filter(User.email == "admin@company.com").one()


def _overdue_task(db, creator: User, assignee: User, days: int) -> Task:
    task = Task(
        tenant_id=creator.tenant_id, title=f"Báo cáo công nợ {uuid.uuid4().hex[:4]}", creator_id=creator.id,
        assignee_id=assignee.id, priority="HIGH", status="PENDING", attachments=[],
        due_date=datetime.now(timezone.utc) - timedelta(days=days),
    )
    db.add(task)
    db.commit()
    return task


# ------------------------------------------------------------------------------ rules


def test_every_rule_has_a_default_and_a_threshold_out_of_range_is_refused(db, admin):
    settings = {setting.rule.key: setting for setting in rule_settings(db, admin.tenant_id)}
    assert set(settings) == {rule.key for rule in RULES}
    assert all(setting.enabled for setting in settings.values())

    with pytest.raises(HTTPException) as refused:
        save_rule_settings(db, admin, [{"key": "APPROVAL_WAITING", "enabled": True, "threshold": 0}])
    assert refused.value.status_code == 422
    with pytest.raises(HTTPException) as unknown:
        save_rule_settings(db, admin, [{"key": "WEATHER", "enabled": True, "threshold": 1}])
    assert unknown.value.status_code == 422

    saved = {s.rule.key: s for s in save_rule_settings(db, admin, [{"key": "TASK_OVERDUE", "enabled": False, "threshold": 4}])}
    assert saved["TASK_OVERDUE"].enabled is False and saved["TASK_OVERDUE"].threshold == 4
    db.rollback()


def test_work_handed_out_and_left_late_raises_an_alert_past_the_threshold(db):
    director = person(db, "executive.briefing.view")
    worker = person(db)
    task = _overdue_task(db, director, worker, days=3)

    found = [alert for alert in evaluate_alerts(db, director) if alert["rule"] == "TASK_OVERDUE"]
    assert [alert["subject"] for alert in found] == [str(task.id)]
    assert "quá hạn 3 ngày" in found[0]["message"] and found[0]["link"] == "/tasks"

    stricter = [
        alerts_module.Setting(setting.rule, setting.enabled, 5 if setting.rule.key == "TASK_OVERDUE" else setting.threshold)
        for setting in rule_settings(db, director.tenant_id)
    ]
    assert not [alert for alert in evaluate_alerts(db, director, stricter) if alert["rule"] == "TASK_OVERDUE"]


def test_a_rule_never_reads_a_department_the_recipient_may_not_see(db, monkeypatch):
    director = person(db, "executive.briefing.view")

    def forbidden(*_args, **_kwargs):
        raise AssertionError("read a department outside the recipient's permissions")

    for name in ("budget_vs_actual", "aging", "list_contract_status_summaries", "get_llm_cost_summary"):
        monkeypatch.setattr(alerts_module, name, forbidden)

    found = evaluate_alerts(db, director)
    assert {alert["rule"] for alert in found} <= {"APPROVAL_WAITING", "TASK_OVERDUE"}


def test_the_digest_skips_what_the_reader_may_not_see():
    briefing = {"sections": {
        "approvals": {"permitted": True, "waiting_count": 4, "by_risk": {"cao": 1, "trung bình": 3}},
        "finance": {"permitted": False, "needs": ["Xem sổ sách kế toán"]},
        "people": {"permitted": True, "on_leave": {"people": 2}, "contracts_ending": {"permitted": False}},
        "tasks": {"permitted": True, "handed_out_overdue": 1},
    }}
    lines = digest_lines(briefing, alert_count=2)
    assert lines == [
        "Chờ bạn duyệt: 4 yêu cầu, 1 rủi ro cao",
        "Cảnh báo đang mở: 2",
        "Nghỉ hôm nay: 2 người",
        "Việc bạn giao đã trễ hạn: 1",
    ]


# -------------------------------------------------------------------------- scheduler


def test_the_digest_goes_out_once_on_a_chosen_day_and_alerts_are_not_repeated(db, monkeypatch):
    director = person(db, "executive.briefing.view")
    _overdue_task(db, director, person(db), days=2)
    monkeypatch.setattr(scheduler, "recipients", lambda _db, _tenant: [director])
    row = scheduler.get_schedule(db, director.tenant_id)
    row.digest_weekdays, row.digest_hour, row.last_digest_on, row.last_alert_check_at = [1], 8, None, None
    db.commit()

    def mine(event_type: str) -> int:
        return db.query(Notification).filter(
            Notification.user_id == director.id, Notification.event_type == event_type
        ).count()

    monday_early = datetime(2032, 3, 15, 7, 30, tzinfo=VN)  # a Monday, before the hour
    done = scheduler.run_tenant(db, director.tenant_id, monday_early)
    assert done["digests"] == 0 and done["alerts"] >= 1
    alerts_sent = mine("EXECUTIVE_ALERT")

    done = scheduler.run_tenant(db, director.tenant_id, monday_early + timedelta(hours=1))
    assert done["digests"] == 1 and mine("EXECUTIVE_DIGEST") == 1
    # Checked again within the interval: not re-run. After it: deduplicated for the day.
    assert mine("EXECUTIVE_ALERT") == alerts_sent
    scheduler.run_tenant(db, director.tenant_id, monday_early + timedelta(hours=3))
    assert mine("EXECUTIVE_ALERT") == alerts_sent

    assert scheduler.run_tenant(db, director.tenant_id, monday_early + timedelta(hours=5))["digests"] == 0
    tuesday = monday_early + timedelta(days=1, hours=2)
    assert scheduler.run_tenant(db, director.tenant_id, tuesday)["digests"] == 0  # not a chosen day
    digest = db.query(Notification).filter(
        Notification.user_id == director.id, Notification.event_type == "EXECUTIVE_DIGEST"
    ).one()
    assert digest.entity_type == "EXECUTIVE" and "Việc bạn giao đã trễ hạn: 1" in digest.message
    db.rollback()


def test_switched_off_the_scheduler_sends_nothing(db, monkeypatch):
    director = person(db, "executive.briefing.view")
    _overdue_task(db, director, person(db), days=2)
    monkeypatch.setattr(scheduler, "recipients", lambda _db, _tenant: [director])
    row = scheduler.get_schedule(db, director.tenant_id)
    row.digest_enabled = row.alerts_enabled = False
    db.commit()

    done = scheduler.run_tenant(db, director.tenant_id, datetime(2032, 3, 15, 10, 0, tzinfo=VN))
    assert done == {"alerts": 0, "digests": 0}
    db.rollback()


def test_the_schedule_refuses_an_hour_or_day_that_does_not_exist(db, admin):
    for change in ({"digest_hour": 24}, {"digest_weekdays": [0, 3]}):
        with pytest.raises(HTTPException) as refused:
            scheduler.save_schedule(db, admin, change)
        assert refused.value.status_code == 422
    db.rollback()


# ------------------------------------------------------------------------------- page


def test_the_page_needs_the_ceo_box_and_settings_need_the_alerts_box(client, db):
    outsider = login(client, person(db))
    assert client.get("/api/v1/executive/briefing", headers=outsider).status_code == 403

    director = person(db, "executive.briefing.view")
    headers = login(client, director)
    for path in ("/api/v1/executive/briefing?period=WEEK", "/api/v1/executive/approvals", "/api/v1/executive/tasks",
                 "/api/v1/executive/alerts", "/api/v1/executive/digest", "/api/v1/executive/people"):
        assert client.get(path, headers=headers).status_code == 200, path
    alerts = client.get("/api/v1/executive/alerts", headers=headers).json()
    assert alerts["can_manage"] is False and len(alerts["rules"]) == len(RULES)
    body = {"rules": [{"key": "TASK_OVERDUE", "enabled": True, "threshold": 2}]}
    assert client.put("/api/v1/executive/alert-settings", json=body, headers=headers).status_code == 403

    manager = login(client, person(db, "executive.briefing.view", "executive.alerts.manage"))
    saved = client.put(
        "/api/v1/executive/alert-settings",
        json={**body, "schedule": {"digest_hour": 7, "digest_weekdays": [1, 3, 5]}},
        headers=manager,
    )
    assert saved.status_code == 200, saved.text
    payload = saved.json()
    assert payload["can_manage"] is True
    assert next(rule for rule in payload["rules"] if rule["key"] == "TASK_OVERDUE")["threshold"] == 2
    assert payload["schedule"]["digest_hour"] == 7 and payload["schedule"]["digest_weekdays"] == [1, 3, 5]


def test_work_is_handed_out_from_the_page_to_someone_in_reach(client, db, ceo_token_headers):
    admin_user = db.query(User).filter(User.email == "admin@company.com").one()
    colleague = db.query(User).filter(User.email == "employee@company.com").one()
    due = (date.today() + timedelta(days=5)).isoformat()

    created = client.post(
        "/api/v1/executive/tasks",
        json={"assignee_id": str(colleague.id), "title": "Chuẩn bị báo cáo quý", "due_date": due, "priority": "HIGH"},
        headers=ceo_token_headers,
    )
    assert created.status_code == 201, created.text
    task = db.get(Task, uuid.UUID(created.json()["task_id"]))
    assert task.creator_id == admin_user.id and task.assignee_id == colleague.id and task.priority == "HIGH"
    assert db.query(Notification).filter(
        Notification.user_id == colleague.id, Notification.dedup_key == f"task-assigned:{task.id}"
    ).count() == 1

    unknown = client.post(
        "/api/v1/executive/tasks", json={"assignee_id": str(uuid.uuid4()), "title": "x"}, headers=ceo_token_headers,
    )
    assert unknown.status_code == 404
    late = client.post(
        "/api/v1/executive/tasks",
        json={"assignee_id": str(colleague.id), "title": "x", "due_date": "2020-01-01"},
        headers=ceo_token_headers,
    )
    assert late.status_code == 422 and "đã qua" in late.json()["detail"]


def test_the_briefing_downloads_as_a_workbook(client, ceo_token_headers):
    from io import BytesIO

    from openpyxl import load_workbook

    response = client.get("/api/v1/executive/briefing/export?period=MONTH", headers=ceo_token_headers)
    assert response.status_code == 200
    book = load_workbook(BytesIO(response.content))
    assert book.sheetnames[0] == "Tóm tắt" and "Chờ duyệt" in book.sheetnames
