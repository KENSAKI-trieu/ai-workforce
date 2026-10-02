"""Proposing journal entries for invoices, and what approving, editing or rejecting does."""

from __future__ import annotations

import io
import uuid
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from app.domains.finance.journal_proposal import propose_for_invoice, ProposalRefused
from app.domains.finance.settings import get_settings, seed_chart
from app.models.models import (
    FinInvoice,
    FinJournalEntry,
    FinLedgerLine,
    FinParty,
    FinPostingRule,
    User,
    WorkflowApproval,
)
from app.tools.executors.finance import propose_journal_entry
from app.tools.registry import ToolContext, tool_registry
from app.tools.schemas import ProposeJournalEntryInput
from tests.finance_helpers import login, person, unique_tax_code


@pytest.fixture()
def books(transactional_db_session):
    db = transactional_db_session
    tenant_id = db.query(User).filter(User.email == "employee@company.com").one().tenant_id
    get_settings(db, tenant_id)
    seed_chart(db, tenant_id, "TT200")
    party = FinParty(tenant_id=tenant_id, kind="VENDOR", tax_code=unique_tax_code(), name="Sao Mai", is_new=True)
    db.add(party)
    db.commit()
    return tenant_id, party


def _invoice(db, tenant_id, party, *, status="MATCHED", total="11000000", vat="1000000", direction="IN") -> FinInvoice:
    total, vat = Decimal(total), Decimal(vat)
    invoice = FinInvoice(
        tenant_id=tenant_id, direction=direction, party_id=party.id,
        seller_tax_code=party.tax_code if direction == "IN" else "", seller_name=party.name,
        series="C26TAA", number=str(uuid.uuid4().int)[:7], issue_date=date(2026, 9, 15),
        amount_before_tax=total - vat, vat_amount=vat, total_amount=total, status=status,
        lines=[{"kind": "1", "name": "Giấy in A4", "amount": str(total - vat)}],
    )
    db.add(invoice)
    db.commit()
    return invoice


def _approval(db, entry: FinJournalEntry) -> WorkflowApproval:
    return db.query(WorkflowApproval).filter(WorkflowApproval.workflow_id == entry.workflow_id).one()


class _Model:
    enabled = True

    def __init__(self, account: str):
        self.account = account
        self.seen: list = []

    def generate_text(self, messages, **kwargs):
        self.seen.append(messages)
        return {"provider": "gemini", "content": f'{{"account": "{self.account}", "reason": "Văn phòng phẩm"}}'}


@pytest.fixture()
def drafter(transactional_db_session):
    return person(transactional_db_session, "finance.journal.draft", "finance.invoice.process")


@pytest.fixture()
def approver(transactional_db_session):
    return person(transactional_db_session, "finance.journal.approve", "finance.journal.approve_high")


def test_a_named_account_drafts_a_balanced_entry_from_the_invoice(books, drafter, transactional_db_session):
    db = transactional_db_session
    tenant_id, party = books
    invoice = _invoice(db, tenant_id, party)
    entry, created = propose_for_invoice(db, drafter, invoice.id, main_account="6423")
    db.commit()
    assert created and entry.status == "PENDING_APPROVAL" and entry.proposed_by == "USER"
    assert [(line.account_code, line.debit, line.credit) for line in entry.lines] == [
        ("6423", Decimal("10000000.00"), Decimal("0")),
        ("1331", Decimal("1000000.00"), Decimal("0")),
        ("331", Decimal("0"), Decimal("11000000.00")),
    ]
    # A new vendor is a reason to look closer even when the user chose the account.
    assert entry.confidence == "LOW"
    approval = _approval(db, entry)
    assert approval.action_type == "FINANCE_JOURNAL_APPROVAL"
    assert approval.payload["required_permission"] == "finance.journal.approve"
    again, created_again = propose_for_invoice(db, drafter, invoice.id, main_account="6427")
    assert not created_again and again.id == entry.id


def test_an_invoice_with_open_exceptions_is_not_proposed(books, drafter, transactional_db_session):
    tenant_id, party = books
    invoice = _invoice(transactional_db_session, tenant_id, party, status="EXCEPTION")
    invoice.exceptions = [{"code": "PO_NOT_FOUND", "message": "Không có PO", "severity": "BLOCKING"}]
    transactional_db_session.commit()
    with pytest.raises(ProposalRefused) as refused:
        propose_for_invoice(transactional_db_session, drafter, invoice.id, main_account="6423")
    assert "Không có PO" in str(refused.value)


def test_the_model_chooses_only_among_the_companys_accounts(books, drafter, transactional_db_session, monkeypatch):
    db = transactional_db_session
    tenant_id, party = books
    model = _Model("6427")
    monkeypatch.setattr("app.domains.finance.journal_proposal.get_ai_service_client", lambda: model)
    entry, _ = propose_for_invoice(db, drafter, _invoice(db, tenant_id, party).id)
    assert entry.proposed_by == "MODEL" and entry.lines[0].account_code == "6427"
    assert entry.confidence == "LOW"
    offered = model.seen[0][1]["content"]
    assert '"6427"' in offered and '"1111"' not in offered  # cash is never a candidate
    monkeypatch.setattr("app.domains.finance.journal_proposal.get_ai_service_client", lambda: _Model("1111"))
    with pytest.raises(ProposalRefused):
        propose_for_invoice(db, drafter, _invoice(db, tenant_id, party).id)


def test_approving_unchanged_posts_and_teaches_the_rule(
    client, books, drafter, approver, transactional_db_session
):
    db = transactional_db_session
    tenant_id, party = books
    invoice = _invoice(db, tenant_id, party)
    entry, _ = propose_for_invoice(db, drafter, invoice.id, main_account="6423")
    db.commit()
    approval = _approval(db, entry)
    # The one who asked cannot sign.
    own = client.post(f"/api/v1/approvals/{approval.id}/action", json={"action": "APPROVE"}, headers=login(client, drafter))
    assert own.status_code == 403
    signed = client.post(f"/api/v1/approvals/{approval.id}/action", json={"action": "APPROVE"}, headers=login(client, approver))
    assert signed.status_code == 200, signed.text
    db.refresh(entry)
    db.refresh(invoice)
    db.refresh(party)
    assert entry.status == "POSTED" and invoice.status == "POSTED" and not party.is_new
    ledger = db.query(FinLedgerLine).filter(FinLedgerLine.journal_entry_id == entry.id).all()
    assert sum(line.debit for line in ledger) == sum(line.credit for line in ledger) == Decimal("11000000.00")
    assert {line.period for line in ledger} == {"2026-09"}
    rule = db.query(FinPostingRule).filter(FinPostingRule.party_id == party.id).one()
    assert rule.main_account == "6423"

    # The next invoice from this vendor follows the rule, with nothing left to doubt.
    following, _ = propose_for_invoice(db, drafter, _invoice(db, tenant_id, party).id)
    assert following.proposed_by == "RULE"
    assert following.lines[0].account_code == "6423"
    assert following.confidence == "HIGH"


def test_an_approvers_correction_is_posted_and_remembered_not_learned(
    client, books, drafter, approver, transactional_db_session
):
    db = transactional_db_session
    tenant_id, party = books
    entry, _ = propose_for_invoice(db, drafter, _invoice(db, tenant_id, party).id, main_account="6423")
    db.commit()
    approval = _approval(db, entry)
    edited = client.post(
        f"/api/v1/approvals/{approval.id}/action",
        json={"action": "EDIT_AND_APPROVE", "edited_payload": {"lines": [{"line_no": 1, "account_code": "6427"}],
              "amount": "1", "required_permission": "finance.journal.approve"}},
        headers=login(client, approver),
    )
    assert edited.status_code == 200, edited.text
    db.refresh(entry)
    db.refresh(approval)
    assert entry.status == "POSTED"
    assert entry.lines[0].account_code == "6427"
    assert entry.corrections["changes"] == [{"line_no": 1, "from": "6423", "to": "6427"}]
    # An edit cannot rewrite what the server decided.
    assert approval.payload["amount"] == "11000000.00"
    assert db.query(FinPostingRule).filter(FinPostingRule.party_id == party.id).first() is None


def test_an_edit_to_a_header_account_is_refused(client, books, drafter, approver, transactional_db_session):
    db = transactional_db_session
    tenant_id, party = books
    entry, _ = propose_for_invoice(db, drafter, _invoice(db, tenant_id, party).id, main_account="6423")
    db.commit()
    response = client.post(
        f"/api/v1/approvals/{_approval(db, entry).id}/action",
        json={"action": "EDIT_AND_APPROVE", "edited_payload": {"lines": [{"line_no": 1, "account_code": "642"}]}},
        headers=login(client, approver),
    )
    assert response.status_code == 422


def test_a_rejected_entry_leaves_the_invoice_to_propose_again(
    client, books, drafter, approver, transactional_db_session
):
    db = transactional_db_session
    tenant_id, party = books
    invoice = _invoice(db, tenant_id, party)
    entry, _ = propose_for_invoice(db, drafter, invoice.id, main_account="6423")
    db.commit()
    rejected = client.post(
        f"/api/v1/approvals/{_approval(db, entry).id}/action",
        json={"action": "REJECT", "comments": "Sai TK"}, headers=login(client, approver),
    )
    assert rejected.status_code == 200, rejected.text
    db.refresh(entry)
    db.refresh(invoice)
    assert entry.status == "REJECTED" and invoice.status == "MATCHED"
    retry, created = propose_for_invoice(db, drafter, invoice.id, main_account="6427")
    assert created and retry.id != entry.id


def test_a_sales_invoice_credits_revenue_and_output_vat(books, drafter, transactional_db_session):
    db = transactional_db_session
    tenant_id, party = books
    invoice = _invoice(db, tenant_id, party, direction="OUT", total="22000000", vat="2000000")
    entry, _ = propose_for_invoice(db, drafter, invoice.id, main_account="5113")
    assert [(line.account_code, line.debit, line.credit) for line in entry.lines] == [
        ("131", Decimal("22000000.00"), Decimal("0")),
        ("5113", Decimal("0"), Decimal("20000000.00")),
        ("33311", Decimal("0"), Decimal("2000000.00")),
    ]
    with pytest.raises(ProposalRefused):
        propose_for_invoice(db, drafter, _invoice(db, tenant_id, party, direction="OUT").id, main_account="6423")


def test_the_tool_answers_with_the_entry_and_is_terminal(books, drafter, transactional_db_session):
    db = transactional_db_session
    tenant_id, party = books
    invoice = _invoice(db, tenant_id, party)
    request = ProposeJournalEntryInput.model_construct(
        tenant_id=tenant_id, audit=None, invoice_id=invoice.id, main_account="6423",
    )
    result = propose_journal_entry(ToolContext(db=db, actor=drafter), request)
    assert result["created"] and "gửi duyệt" in result["reply"]
    assert "Nợ 6423 10.000.000 ₫" in result["reply"]
    definition = tool_registry.get("propose_journal_entry")
    assert definition.terminal and definition.opens_approval
    assert not definition.acl.permits(person(db, "finance.invoice.process"))


def test_posted_entries_export_as_debit_credit_pairs(
    client, books, drafter, approver, transactional_db_session
):
    db = transactional_db_session
    tenant_id, party = books
    entry, _ = propose_for_invoice(db, drafter, _invoice(db, tenant_id, party).id, main_account="6423")
    db.commit()
    client.post(
        f"/api/v1/approvals/{_approval(db, entry).id}/action", json={"action": "APPROVE"}, headers=login(client, approver),
    ).raise_for_status()
    reader = person(db, "finance.ledger.view")
    response = client.get("/api/v1/finance/export/journal.xlsx?period=2026-09", headers=login(client, reader))
    assert response.status_code == 200
    rows = list(load_workbook(io.BytesIO(response.content)).active.iter_rows(values_only=True))
    mine = [row for row in rows[1:] if row[2] == entry.description]
    assert sorted((row[3], row[4], row[5]) for row in mine) == [
        ("1331", "331", 1000000.0), ("6423", "331", 10000000.0),
    ]
