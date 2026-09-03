"""
Tests for HR Agent intent handling, tool execution, leave balances, and Approval Cards.
"""

import json
from datetime import date, timedelta
from io import BytesIO

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from pypdf import PdfReader

from app.core.database import SyncSessionLocal
from app.models.models import AIAgent, AuditLog, User, UserMemory
from app.services.agents.agent_executor import (
    HR_CONFIGURATION_VERSION,
    HR_RETIRED_TOOLS,
    _classify_hr_intent,
    _leave_balance_names_another_person,
    _repair_hr_agent_capabilities,
    _require_tool,
)


@pytest.mark.parametrize(
    ("message", "expected"),
    [
        ("tìm các nhân viên quản lý", "MANAGER_DIRECTORY"),
        ("tim cac nhan vien quan ly", "MANAGER_DIRECTORY"),
        ("có bao nhiêu quản lý", "MANAGER_DIRECTORY"),
        ("công ty có mấy nhân viên", "EMPLOYEE_DIRECTORY"),
        ("số lượng nhân sự", "EMPLOYEE_DIRECTORY"),
        ("danh sach nhan vien", "EMPLOYEE_DIRECTORY"),
        ("luong cua toi", "SELF_COMPENSATION"),
        ("hop dong sap het han", "CONTRACT_EXPIRY"),
        ("quy định nghỉ phép là gì", "POLICY_QUERY"),
        ("có bao nhiêu nhân viên đang nghỉ phép", "EMPLOYEE_LEAVE_STATUS_COUNT"),
        ("xin chào", "UNKNOWN"),
    ],
)
def test_hr_intent_classifier_normalizes_actions_and_entities(message, expected):
    assert _classify_hr_intent(message) == expected


def test_stale_hr_agent_capabilities_are_split_into_narrow_profile_tools(
    client,
    employee_token_headers,
    transactional_db_session,
):
    agent = transactional_db_session.query(AIAgent).filter(
        AIAgent.role_code == "HR"
    ).first()
    agent.tools_access = [
        "query_leave_balance",
        "request_leave",
        "hybrid_rag_search",
        "get_employee_profile",
    ]
    agent.allowed_actions = list(agent.tools_access)
    agent.disallowed_actions = []
    agent.configuration_version = 1
    _repair_hr_agent_capabilities(agent)
    transactional_db_session.commit()

    assert "get_employee_profile" not in agent.tools_access
    assert "get_employee_full_profile" in agent.tools_access
    # Directory access follows profile access. This used to key off the basic-profile
    # grant, which version 7 retired; the outcome for a version 1 row is unchanged.
    assert "query_company_users_sql" in agent.tools_access
    assert "export_hr_directory" in agent.tools_access
    # Tools no branch dispatches and the tool gateway does not define are revoked, not
    # left as grants with nothing behind them.
    assert not HR_RETIRED_TOOLS & set(agent.tools_access)
    assert not HR_RETIRED_TOOLS & set(agent.allowed_actions)
    assert not HR_RETIRED_TOOLS & set(agent.disallowed_actions)
    assert agent.configuration_version == HR_CONFIGURATION_VERSION

    response = client.post(
        "/api/v1/agent/chat",
        json={"agent_role": "HR", "message": "Hồ sơ của tôi"},
        headers=employee_token_headers,
    )
    assert response.status_code == 200, response.text
    assert response.json()["hr_card"]["type"] == "EMPLOYEE_PROFILE"


def test_version_seven_revokes_the_unreachable_basic_profile_grant(
    transactional_db_session,
):
    """An agent already stamped at version 6 still carries the name; 7 has to sweep it."""
    agent = transactional_db_session.query(AIAgent).filter(
        AIAgent.role_code == "HR"
    ).first()
    # Set the grants explicitly: this module shares one session, so whatever an earlier
    # test left on the row must not decide what this one is starting from.
    agent.tools_access = [
        "get_employee_basic_profile",
        "get_employee_full_profile",
        "query_leave_balance",
    ]
    agent.allowed_actions = list(agent.tools_access)
    agent.disallowed_actions = []
    agent.configuration_version = 6

    _repair_hr_agent_capabilities(agent)

    assert "get_employee_basic_profile" not in agent.tools_access
    assert "get_employee_basic_profile" not in agent.allowed_actions
    assert agent.configuration_version == HR_CONFIGURATION_VERSION
    # The sweep must not cost the agent a capability it can still dispatch.
    assert "get_employee_full_profile" in agent.tools_access


def test_a_denied_profile_grant_still_denies_the_directory_tool():
    """Version 4 keyed off the basic-profile name until 7 retired it; both legacy paths
    have to keep reaching the same verdict."""
    class _Agent:
        role_code = "HR"
        configuration_version = 1
        tools_access = ["query_leave_balance", "get_employee_profile"]
        allowed_actions = ["query_leave_balance", "get_employee_profile"]
        disallowed_actions = ["get_employee_profile"]

    agent = _Agent()
    _repair_hr_agent_capabilities(agent)

    # The grant survives in tools_access because version 2 hands out the whole core set;
    # the explicit denial is what has to win, and _require_tool checks it first.
    assert "query_company_users_sql" in agent.disallowed_actions
    with pytest.raises(HTTPException) as denied:
        _require_tool(agent, "query_company_users_sql")
    assert denied.value.status_code == 403


def test_hr_company_user_sql_tool_respects_chat_actor_scope(
    client,
    employee_token_headers,
    ceo_token_headers,
):
    employee_response = client.post(
        "/api/v1/agent/chat",
        json={"agent_role": "HR", "message": "Danh sách nhân viên"},
        headers=employee_token_headers,
    )
    assert employee_response.status_code == 200, employee_response.text
    employee_data = employee_response.json()
    assert employee_data["tools_executed"][0]["tool_name"] == "query_company_users_sql"
    assert employee_data["hr_card"]["scope"] == "SELF"
    assert len(employee_data["hr_card"]["items"]) == 1
    assert employee_data["hr_card"]["items"][0]["employee"]["email"] == "employee@company.com"

    ceo_response = client.post(
        "/api/v1/agent/chat",
        json={"agent_role": "HR", "message": "Danh sách nhân viên"},
        headers=ceo_token_headers,
    )
    assert ceo_response.status_code == 200, ceo_response.text
    ceo_data = ceo_response.json()
    assert ceo_data["tools_executed"][0]["tool_name"] == "query_company_users_sql"
    assert ceo_data["hr_card"]["scope"] == "COMPANY"
    assert len(ceo_data["hr_card"]["items"]) > 1


def test_hr_manager_directory_phrase_from_chat_routes_to_sql(
    client,
    ceo_token_headers,
):
    response = client.post(
        "/api/v1/agent/chat",
        json={"agent_role": "HR", "message": "tìm các nhân viên quản lý"},
        headers=ceo_token_headers,
    )
    assert response.status_code == 200, response.text
    data = response.json()
    assert data["tools_executed"][0]["tool_name"] == "query_company_users_sql"
    assert data["hr_card"]["directory_type"] == "MANAGERS"
    assert data["hr_card"]["total_count"] == len(data["hr_card"]["items"])
    assert data["hr_card"]["items"]
    assert {
        item["employee"]["role"] for item in data["hr_card"]["items"]
    } <= {"Admin", "Manager"}


def test_hr_unknown_query_does_not_fall_back_to_policy(
    client,
    ceo_token_headers,
):
    unknown = client.post(
        "/api/v1/agent/chat",
        json={"agent_role": "HR", "message": "xin chào"},
        headers=ceo_token_headers,
    )
    assert unknown.status_code == 200, unknown.text
    assert unknown.json()["tools_executed"] == []
    assert "chưa xác định rõ" in unknown.json()["reply"].lower()


def test_hr_unsupported_leave_statistics_states_the_gap_then_still_searches(
    client,
    ceo_token_headers,
):
    """There is no day-by-day leave calendar tool, but the question is still answerable.

    Both the keyword rules and the LLM router can land on this intent, and neither has
    an alternative label to fall back to, so the branch must not end the turn empty.
    """
    unsupported = client.post(
        "/api/v1/agent/chat",
        json={"agent_role": "HR", "message": "có bao nhiêu nhân viên đang nghỉ phép"},
        headers=ceo_token_headers,
    )
    assert unsupported.status_code == 200, unsupported.text
    data = unsupported.json()
    assert "chưa có tool" in data["reply"].lower()
    assert [item["tool_name"] for item in data["tools_executed"]] == ["hybrid_rag_search"]


def test_hr_export_intent_requires_scope_and_format(
    client,
    ceo_token_headers,
):
    for message in ("Xuất file", "Trích xuất file"):
        incomplete = client.post(
            "/api/v1/agent/chat",
            json={"agent_role": "HR", "message": message},
            headers=ceo_token_headers,
        )
        assert incomplete.status_code == 200, incomplete.text
        incomplete_data = incomplete.json()
        assert incomplete_data["tools_executed"] == []
        assert "loại dữ liệu" in incomplete_data["reply"]
        assert "định dạng" in incomplete_data["reply"]

    ready = client.post(
        "/api/v1/agent/chat",
        json={"agent_role": "HR", "message": "Xuất danh sách nhân viên Excel"},
        headers=ceo_token_headers,
    )
    assert ready.status_code == 200, ready.text
    ready_data = ready.json()
    assert ready_data["tools_executed"][0]["tool_name"] == "export_hr_directory"
    assert ready_data["hr_card"]["type"] == "FILE_EXPORT"
    assert ready_data["hr_card"]["format"] == "xlsx"
    assert ready_data["hr_card"]["directory_type"] == "employees"
    assert ready_data["hr_card"]["download_url"].startswith(
        "/api/v1/hr/employees/export?"
    )


def test_hr_directory_export_formats_and_scope(
    client,
    employee_token_headers,
    ceo_token_headers,
    transactional_db_session,
):
    excel = client.get(
        "/api/v1/hr/employees/export?format=xlsx&directory=employees",
        headers=ceo_token_headers,
    )
    assert excel.status_code == 200, excel.text
    assert excel.content.startswith(b"PK")
    assert excel.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    workbook = load_workbook(BytesIO(excel.content), read_only=True)
    sheet = workbook["Danh sach nhan su"]
    assert [cell.value for cell in sheet[4]] == [
        "Mã nhân viên",
        "Họ và tên",
        "Email",
        "Vai trò",
        "Phòng ban",
        "Chức danh",
        "Trạng thái",
        "Quản lý trực tiếp",
    ]
    assert sheet.max_row > 4

    pdf = client.get(
        "/api/v1/hr/employees/export?format=pdf&directory=managers",
        headers=ceo_token_headers,
    )
    assert pdf.status_code == 200, pdf.text
    assert pdf.content.startswith(b"%PDF-")
    assert pdf.headers["content-type"].startswith("application/pdf")
    pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(pdf.content)).pages)
    assert "DANH SÁCH NHÂN SỰ" in pdf_text

    own_json = client.get(
        "/api/v1/hr/employees/export?format=json&directory=employees",
        headers=employee_token_headers,
    )
    assert own_json.status_code == 200, own_json.text
    payload = own_json.json()
    assert payload["metadata"]["scope"] == "SELF"
    assert payload["metadata"]["total_count"] == 1
    assert [item["email"] for item in payload["items"]] == ["employee@company.com"]
    assert all("monthly_salary" not in item for item in payload["items"])

    audit = transactional_db_session.query(AuditLog).filter(
        AuditLog.tool_name == "export_hr_directory"
    ).order_by(AuditLog.created_at.desc()).first()
    assert audit is not None
    assert audit.output_result["allowed_sections"] == ["BASIC"]


def test_hr_leave_balance_query(client, employee_token_headers):
    """Test HR Agent handling leave balance inquiry."""
    payload = {
        "agent_role": "HR",
        "message": "Tôi còn bao nhiêu ngày phép?",
    }
    response = client.post("/api/v1/agent/chat", json=payload, headers=employee_token_headers)
    assert response.status_code == 200
    data = response.json()
    assert data["agent_role"] == "HR"
    assert "ngày phép" in data["reply"].lower()
    assert len(data["tools_executed"]) > 0
    assert data["tools_executed"][0]["tool_name"] == "query_leave_balance"
    assert data["hr_card"]["type"] == "LEAVE_BALANCE"

    conversation = client.get(
        f"/api/v1/agent/conversations/{data['conversation_id']}",
        headers=employee_token_headers,
    )
    assert conversation.status_code == 200
    assistant = conversation.json()["messages"][-1]
    assert assistant["attachments"][0]["type"] == "HR_CARD"
    assert assistant["attachments"][0]["payload"]["type"] == "LEAVE_BALANCE"


def test_hr_leave_request_creates_approval_card(client, employee_token_headers):
    """Test HR Agent handling leave submission and generating an Approval Card."""
    # Reset leave balance memory for employee to ensure deterministic quota
    db = SyncSessionLocal()
    try:
        user = db.query(User).filter(User.email == "employee@company.com").first()
        if user:
            mem = db.query(UserMemory).filter(
                UserMemory.user_id == user.id,
                UserMemory.memory_key == "leave_balance",
            ).first()
            if mem:
                mem.memory_value = json.dumps({"total_days": 12, "used_days": 2, "remaining_days": 10})
                db.commit()
    finally:
        db.close()

    start = date.today() + timedelta(days=30)
    while start.weekday() >= 5:
        start += timedelta(days=1)
    end = start + timedelta(days=1)
    while end.weekday() >= 5:
        end += timedelta(days=1)
    payload = {
        "agent_role": "HR",
        "message": (
            f"Tôi muốn xin nghỉ phép từ {start.isoformat()} đến {end.isoformat()} "
            "vì lý do gia đình"
        ),
    }
    response = client.post("/api/v1/agent/chat", json=payload, headers=employee_token_headers)
    assert response.status_code == 200
    data = response.json()
    assert data["agent_role"] == "HR"
    assert data["approval_card"] is not None
    card = data["approval_card"]
    assert card["action_type"] == "XIN NGHỈ PHÉP"
    assert card["status"] == "WAITING"
    assert "id" in card


@pytest.mark.parametrize(
    ("message", "expected_other"),
    [
        ("Tôi còn bao nhiêu ngày phép?", False),
        ("còn bao nhiêu ngày phép", False),
        ("Cho tôi biết số ngày phép", False),
        ("Kiểm tra quỹ phép của tôi", False),
        ("An còn bao nhiêu ngày phép?", True),
        ("Nguyễn Văn A còn bao nhiêu ngày phép", True),
        ("quỹ phép của an.nguyen@company.com", True),
    ],
)
def test_leave_balance_subject_detection(message, expected_other):
    assert _leave_balance_names_another_person(message) is expected_other


def test_hr_leave_balance_refuses_a_question_about_somebody_else(
    client,
    employee_token_headers,
):
    """query_leave_balance only reads the asker's own quota, so it must not answer
    a question about a colleague with the asker's figures under the colleague's name."""
    response = client.post(
        "/api/v1/agent/chat",
        json={"agent_role": "HR", "message": "An còn bao nhiêu ngày phép?"},
        headers=employee_token_headers,
    )

    assert response.status_code == 200, response.text
    data = response.json()
    assert data["tools_executed"] == []
    assert data["hr_card"] is None
    assert "chính bạn" in data["reply"]


def test_hr_export_checks_the_grant_before_asking_for_a_format(
    client,
    ceo_token_headers,
    transactional_db_session,
):
    """H4: permission first, conversation second — as in every other HR branch."""
    actor = transactional_db_session.query(User).filter(
        User.email == "admin@company.com"
    ).one()
    agent = transactional_db_session.query(AIAgent).filter(
        AIAgent.tenant_id == actor.tenant_id,
        AIAgent.role_code == "HR",
    ).one()
    agent.disallowed_actions = sorted(
        set(agent.disallowed_actions or []) | {"export_hr_directory"}
    )
    transactional_db_session.commit()

    response = client.post(
        "/api/v1/agent/chat",
        json={"agent_role": "HR", "message": "Xuất file"},
        headers=ceo_token_headers,
    )

    assert response.status_code == 403, response.text
    assert "export_hr_directory" in response.json()["detail"]


@pytest.mark.parametrize(
    "message",
    [
        # Vietnamese puts the subject on either side of the phrase.
        "quỹ phép của An",
        "quỹ phép của nhân viên An",
        "Cho tôi biết quỹ phép của bạn An",
        "số ngày phép của An là bao nhiêu",
        "số ngày phép còn lại của chị Lan",
        "xem quỹ phép của team tôi",
    ],
)
def test_leave_balance_detects_a_subject_after_the_phrase(message):
    assert _leave_balance_names_another_person(message) is True


@pytest.mark.parametrize(
    "message",
    [
        "quỹ phép của tôi",
        "số ngày phép của mình",
        "quỹ phép của em",
        "cho em hỏi số ngày phép",
    ],
)
def test_leave_balance_first_person_possessives_stay_self_service(message):
    assert _leave_balance_names_another_person(message) is False


def test_hr_leave_balance_refuses_a_possessive_question_about_a_colleague(
    client,
    employee_token_headers,
):
    data = client.post(
        "/api/v1/agent/chat",
        json={"agent_role": "HR", "message": "quỹ phép của An"},
        headers=employee_token_headers,
    ).json()

    assert data["tools_executed"] == []
    assert "chính bạn" in data["reply"]
